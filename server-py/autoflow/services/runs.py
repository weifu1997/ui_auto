"""Run spec resolution, enqueue/execute lifecycle, deletion, and metrics."""
from __future__ import annotations

import json as _json
import uuid
import re
import shutil
from typing import Any
from ..core import (
    days_ago_iso,
    json,
    now,
    parse_json,
    public_flow_output_names,
    safe_artifact_name,
)
from ..resources import as_record
from ._shared import (
    _TERMINAL_RUN_STATUSES,
)


class RunServices:
    """Run spec resolution, enqueue/execute lifecycle, deletion, and metrics."""

    def append_run_event(
        self, run_id: str, kind: str, data: dict[str, Any]
    ) -> None:
        self.database.execute(
            """
            INSERT INTO platform_run_events (run_id, kind, data, created_at)
            VALUES (?, ?, ?, ?)
            """,
            (run_id, kind, json(data), now()),
        )

    def resolve_run_spec(self, input: dict[str, Any]) -> dict[str, Any]:
        from ..http import PlatformError

        project_id = input["projectId"]
        revision = self.published_revision_for(
            project_id,
            input.get("revisionId") or None,
            flow_id=input.get("flowId") or None,
            environment_id=input.get("environmentId") or None,
            allow_superseded=bool(input.get("allowSuperseded")),
        )
        environment = parse_json(revision["environment_snapshot"], {})
        environment_id = input.get("environmentId")
        if not environment_id:
            environment_id = (
                environment.get("id") if isinstance(environment, dict) else ""
            )
        if not environment_id:
            raise PlatformError(400, "ENVIRONMENT_REQUIRED")
        if environment.get("id") != environment_id:
            raise PlatformError(409, "REVISION_ENVIRONMENT_MISMATCH")
        self.require_chromium_environment(
            environment if isinstance(environment, dict) else {}
        )
        self.ensure_chromium_available()
        agent = self.managed_agent(project_id)
        dataset_version_id = input.get("datasetVersionId")
        if not dataset_version_id:
            snapshot_dataset = parse_json(revision["dataset_snapshot"], None)
            if isinstance(snapshot_dataset, dict):
                version_id = snapshot_dataset.get("versionId")
                if isinstance(version_id, str):
                    dataset_version_id = version_id
        dataset_version = (
            self.dataset_version_for(project_id, dataset_version_id)
            if dataset_version_id
            else None
        )
        rows = (
            self.dataset_rows_for(dataset_version["id"])
            if dataset_version
            else [{"rowNumber": None, "data": None}]
        )
        max_runs = input.get("maxRuns")
        if max_runs is not None and len(rows) > int(max_runs):
            raise PlatformError(413, "RUN_COUNT_LIMIT_EXCEEDED")
        flow = parse_json(revision["flow_snapshot"], {})
        flow_steps = flow.get("steps", []) if isinstance(flow, dict) else []
        if not isinstance(flow_steps, list):
            flow_steps = []
        up_to_step_id = input.get("upToStepId")
        if up_to_step_id and not any(
            as_record(step).get("id") == up_to_step_id for step in flow_steps
        ):
            raise PlatformError(400, "RUN_STEP_NOT_FOUND")
        secret_names = flow.get("secretNames", [])
        if not isinstance(secret_names, list):
            secret_names = []
        secret_names = [value for value in secret_names if isinstance(value, str)]
        step_limit = (
            next(
                (
                    index + 1
                    for index, step in enumerate(flow_steps)
                    if as_record(step).get("id") == up_to_step_id
                ),
                0,
            )
            if up_to_step_id
            else len(flow_steps)
        )
        required_secret_names: set[str] = set()
        for step in flow_steps[:step_limit]:
            value = as_record(step).get("value")
            if not isinstance(value, str):
                value = ""
            for name in secret_names:
                if (
                    f"{{{{{name}}}}}" in value
                    or f"{{{{ {name} }}}}" in value
                    or f"{{{{secret.{name}}}}}" in value
                    or f"{{{{ secret.{name} }}}}" in value
                ):
                    required_secret_names.add(name)
        return {
            "projectId": project_id,
            "revision": revision,
            "environment": environment,
            "environmentId": environment_id,
            "datasetVersionId": dataset_version_id,
            "datasetVersion": dataset_version,
            "rows": rows,
            "flow": flow,
            "flowSteps": flow_steps,
            "secretNames": secret_names,
            "requiredSecretNames": required_secret_names,
            "upToStepId": up_to_step_id or None,
            "agent": agent,
        }

    @staticmethod
    def _required_retry_variable_names(
        flow: dict[str, Any], secret_names: list[str]
    ) -> set[str]:
        secret_set = {str(name) for name in secret_names}
        required: set[str] = set()
        steps = flow.get("steps") if isinstance(flow, dict) else []
        if not isinstance(steps, list):
            steps = []
        for step in steps:
            value = as_record(step).get("value")
            if not isinstance(value, str):
                continue
            for match in re.finditer(r"{{\s*([^}]+)\s*}}", value):
                expression = match.group(1).strip()
                if not expression:
                    continue
                if (
                    expression.startswith("secret.")
                    or expression in secret_set
                    or expression.startswith("data.")
                    or expression.startswith("flow.")
                    or expression.startswith("run.")
                    or expression == "env.baseUrl"
                ):
                    continue
                required.add(expression)
        return required

    def _preflight_retry_variables(
        self, project_id: str, flow: dict[str, Any], secret_names: list[str]
    ) -> None:
        from ..http import PlatformError

        required = self._required_retry_variable_names(flow, secret_names)
        if not required:
            return
        rows = self.database.execute(
            """
            SELECT data FROM project_resources
            WHERE project_id = ? AND resource_type = 'variables'
              AND archived_at IS NULL
            """,
            (project_id,),
        ).fetchall()
        available: set[str] = set()
        for variable_row in rows:
            variable = parse_json(variable_row[0], {})
            if (
                not isinstance(variable, dict)
                or variable.get("secret") is True
                or not isinstance(variable.get("name"), str)
                or not isinstance(variable.get("value"), str)
            ):
                continue
            scope = (
                "env"
                if variable.get("scope") == "环境"
                else "project"
                if variable.get("scope") == "项目"
                else ""
            )
            key = f"{scope}.{variable['name']}" if scope else variable["name"]
            available.add(key)
        missing = sorted(required - available)
        if missing:
            raise PlatformError(409, "RUN_VARIABLE_NOT_CONFIGURED")

    def _preflight_retry_spec(self, project_id: str, spec: dict[str, Any]) -> None:
        from ..http import PlatformError

        snapshot_secret_names = [
            name
            for name in spec["secretNames"]
            if name in spec["requiredSecretNames"]
        ]
        missing = self.missing_secret_names(project_id, snapshot_secret_names)
        if missing:
            raise PlatformError(409, "RUN_SECRET_NOT_CONFIGURED")
        self._preflight_retry_variables(
            project_id, spec["flow"], spec["secretNames"]
        )

    @staticmethod
    def _row_from_retry_snapshot(snapshot: dict[str, Any]) -> dict[str, Any]:
        dataset_row = snapshot.get("datasetRow")
        if isinstance(dataset_row, dict):
            return {
                "rowNumber": dataset_row.get("number"),
                "data": dataset_row.get("data"),
            }
        return {"rowNumber": None, "data": None}

    def clone_retry_spec_from_run(self, run: dict[str, Any]) -> dict[str, Any]:
        """从源 run 的持久 snapshot 构造一对一 retry spec，不重新解析 revision/dataset。"""
        from ..http import PlatformError

        snapshot = run.get("snapshot")
        if not isinstance(snapshot, dict):
            raise PlatformError(409, "RUN_SNAPSHOT_NOT_RETRYABLE")
        flow = snapshot.get("flow")
        environment = snapshot.get("environment")
        if not isinstance(flow, dict) or not isinstance(environment, dict):
            raise PlatformError(409, "RUN_SNAPSHOT_NOT_RETRYABLE")
        revision_id = snapshot.get("flowRevisionId")
        revision_checksum = snapshot.get("flowRevisionChecksum")
        if (
            not isinstance(revision_id, str)
            or not isinstance(revision_checksum, str)
            or run.get("revisionId") != revision_id
        ):
            raise PlatformError(409, "RUN_SNAPSHOT_NOT_RETRYABLE")
        if environment.get("id") != run.get("environmentId"):
            raise PlatformError(409, "RUN_SNAPSHOT_NOT_RETRYABLE")
        up_to_step_id = snapshot.get("upToStepId")
        steps = flow.get("steps") if isinstance(flow.get("steps"), list) else []
        if up_to_step_id and not any(
            as_record(step).get("id") == up_to_step_id for step in steps
        ):
            raise PlatformError(409, "RUN_SNAPSHOT_NOT_RETRYABLE")
        elements = snapshot.get("elements")
        if not isinstance(elements, list):
            raise PlatformError(409, "RUN_SNAPSHOT_NOT_RETRYABLE")
        dataset = snapshot.get("dataset")
        if dataset is not None and (
            not isinstance(dataset, dict)
            or not all(
                key in dataset
                for key in ("datasetId", "versionId", "versionNumber", "checksum", "columns")
            )
        ):
            raise PlatformError(409, "RUN_SNAPSHOT_NOT_RETRYABLE")
        dataset_row = snapshot.get("datasetRow")
        if dataset_row is not None and (
            not isinstance(dataset_row, dict)
            or "number" not in dataset_row
            or "data" not in dataset_row
        ):
            raise PlatformError(409, "RUN_SNAPSHOT_NOT_RETRYABLE")
        secret_names = snapshot.get("secretNames", [])
        if not isinstance(secret_names, list) or not all(
            isinstance(name, str) for name in secret_names
        ):
            raise PlatformError(409, "RUN_SNAPSHOT_NOT_RETRYABLE")
        return {
            "projectId": run["projectId"],
            "revision": {
                "id": revision_id,
                "checksum": revision_checksum,
                "element_snapshot": json(elements),
            },
            "environmentId": run["environmentId"],
            "environment": _json.loads(_json.dumps(environment)),
            "flow": _json.loads(_json.dumps(flow)),
            "flowSteps": _json.loads(_json.dumps(steps)),
            "datasetVersionId": (
                dataset.get("versionId") if isinstance(dataset, dict) else None
            ),
            "datasetVersion": (
                {
                    "id": dataset["versionId"],
                    "datasetId": dataset["datasetId"],
                    "versionNumber": dataset["versionNumber"],
                    "checksum": dataset["checksum"],
                    "columns": dataset["columns"],
                }
                if isinstance(dataset, dict)
                else None
            ),
            "secretNames": list(secret_names),
            "requiredSecretNames": set(secret_names),
            "upToStepId": up_to_step_id or None,
            "agent": self.managed_agent(run["projectId"]),
        }

    def retry_run_snapshot(
        self,
        project_id: str,
        run_id: str,
        actor_id: str,
        dispatch_key: str | None = None,
    ) -> dict[str, Any]:
        from ..http import PlatformError

        run = self.run_by_id(run_id)
        if run["projectId"] != project_id:
            raise PlatformError(404, "RUN_NOT_FOUND")
        if run["status"] not in ("failed", "canceled"):
            raise PlatformError(409, "RUN_NOT_RETRYABLE")
        spec = self.clone_retry_spec_from_run(run)
        self._preflight_retry_spec(project_id, spec)
        row = self._row_from_retry_snapshot(run["snapshot"])
        new_run_id: str | None = None
        self.database.execute("BEGIN IMMEDIATE")
        try:
            new_run_id = self.insert_run_from_spec(
                spec,
                row=row,
                created_by=actor_id,
                source="manual",
                dispatch_key=dispatch_key,
                retry_of_run_id=run_id,
            )
            self.database.execute("COMMIT")
        except Exception:
            self.database.execute("ROLLBACK")
            raise
        if new_run_id is None:
            raise RuntimeError("retry_run_snapshot failed to create run")
        self.enqueue_managed_run(new_run_id)
        return {
            "runIds": [new_run_id],
            "runs": [self.run_response(self.run_by_id(new_run_id))],
        }

    def insert_run_from_spec(
        self,
        spec: dict[str, Any],
        *,
        row: dict[str, Any],
        created_by: str,
        source: str,
        dispatch_key: str | None = None,
        batch_id: str | None = None,
        batch_item_index: int | None = None,
        retry_of_run_id: str | None = None,
    ) -> str:
        from ..http import PlatformError

        project_id = spec["projectId"]
        revision = spec["revision"]
        environment_id = spec["environmentId"]
        environment = spec["environment"]
        dataset_version = spec["datasetVersion"]
        dataset_version_id = spec["datasetVersionId"]
        flow = spec["flow"]
        agent = spec["agent"]
        up_to_step_id = spec["upToStepId"]
        if dispatch_key:
            existing = self.database.execute(
                "SELECT id FROM platform_runs WHERE dispatch_key = ? AND project_id = ?",
                (dispatch_key, project_id),
            ).fetchone()
            if existing:
                return existing[0]
        snapshot_secret_names = [
            name
            for name in spec["secretNames"]
            if name in spec["requiredSecretNames"]
        ]
        missing = self.missing_secret_names(project_id, snapshot_secret_names)
        if missing:
            raise PlatformError(409, "RUN_SECRET_NOT_CONFIGURED")
        run_id = str(uuid.uuid4())
        created_at = now()
        # 运行 snapshot 不保存普通变量值；变量在 enqueue/恢复时按项目当前值读取。
        flow_snapshot = dict(flow)
        flow_snapshot.pop("variables", None)
        snapshot = {
            "flowRevisionId": revision["id"],
            "flowRevisionChecksum": revision["checksum"],
            "environmentId": environment_id,
            "flow": flow_snapshot,
            "environment": environment,
            "elements": parse_json(revision["element_snapshot"], []),
            "dataset": (
                {
                    "datasetId": dataset_version["datasetId"],
                    "versionId": dataset_version["id"],
                    "versionNumber": dataset_version["versionNumber"],
                    "checksum": dataset_version["checksum"],
                    "columns": dataset_version["columns"],
                }
                if dataset_version
                else None
            ),
            "datasetRow": (
                {"number": row["rowNumber"], "data": row["data"]}
                if row["data"] is not None
                else None
            ),
            "secretNames": snapshot_secret_names,
            "upToStepId": up_to_step_id,
            "executor": {
                "type": "managed",
                "id": agent["id"],
                "name": agent["name"],
                "browserVersion": agent["browserVersion"],
            },
            "trigger": source,
        }
        if batch_id:
            snapshot["batchId"] = batch_id
            snapshot["batchItemIndex"] = batch_item_index
        self.database.execute(
            """
            INSERT INTO platform_runs (
              id, project_id, revision_id, environment_id, agent_id,
              executor_type, dispatch_key, status, snapshot,
              batch_id, batch_item_index, retry_of_run_id,
              created_by, created_at, updated_at
            ) VALUES (?, ?, ?, ?, ?, 'managed', ?, 'queued', ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                run_id,
                project_id,
                revision["id"],
                environment_id,
                agent["id"],
                dispatch_key,
                json(snapshot),
                batch_id,
                batch_item_index,
                retry_of_run_id,
                created_by,
                created_at,
                created_at,
            ),
        )
        self.append_run_event(
            run_id,
            "run.queued",
            {
                "revisionId": revision["id"],
                "environmentId": environment_id,
                "executorType": "managed",
                "source": source,
                "datasetVersionId": dataset_version_id,
                "datasetRow": row["rowNumber"],
            },
        )
        if retry_of_run_id:
            self.append_run_event(
                run_id,
                "run.retried",
                {"priorRunId": retry_of_run_id, "actorId": created_by},
            )
        return run_id

    def queue_published_runs(self, input: dict[str, Any]) -> dict[str, Any]:
        spec = self.resolve_run_spec(input)
        run_ids: list[str] = []
        self.database.execute("BEGIN IMMEDIATE")
        try:
            for row in spec["rows"]:
                dispatch_key = (
                    f"{input['dispatchKey']}:{row['rowNumber'] or 0}"
                    if input.get("dispatchKey")
                    else None
                )
                run_ids.append(
                    self.insert_run_from_spec(
                        spec,
                        row=row,
                        created_by=input["createdBy"],
                        source=input["source"],
                        dispatch_key=dispatch_key,
                    )
                )
            self.database.execute("COMMIT")
        except Exception:
            self.database.execute("ROLLBACK")
            raise
        for run_id in run_ids:
            self.enqueue_managed_run(run_id)
        return {
            "runIds": run_ids,
            "revision": spec["revision"],
            "environmentId": spec["environmentId"],
            "datasetVersionId": spec["datasetVersionId"],
        }

    def run_by_id(
        self, run_id: str, project_id: str | None = None
    ) -> dict[str, Any]:
        from ..http import PlatformError

        row = self.database.execute(
            """
            SELECT id, project_id, revision_id, environment_id, agent_id,
                   executor_type, status, snapshot, cancellation_requested,
                   result, created_at, updated_at, retry_of_run_id
            FROM platform_runs
            WHERE id = ? AND (? IS NULL OR project_id = ?)
            """,
            (run_id, project_id, project_id),
        ).fetchone()
        if not row:
            raise PlatformError(404, "RUN_NOT_FOUND")
        return {
            "id": row[0],
            "projectId": row[1],
            "revisionId": row[2],
            "environmentId": row[3],
            "agentId": row[4],
            "executorType": row[5],
            "status": row[6],
            "snapshot": parse_json(row[7], {}),
            "cancellationRequested": bool(row[8]),
            "result": parse_json(row[9], None),
            "createdAt": row[10],
            "updatedAt": row[11],
            "retryOfRunId": row[12],
        }

    def delete_run(self, project_id: str, run_id: str) -> dict[str, Any]:
        from ..http import PlatformError

        run = self.database.execute(
            "SELECT id, project_id, status FROM platform_runs WHERE id = ?",
            (run_id,),
        ).fetchone()
        if not run:
            raise PlatformError(404, "RUN_NOT_FOUND")
        if run[1] != project_id:
            raise PlatformError(404, "RUN_NOT_FOUND")
        if run[2] not in _TERMINAL_RUN_STATUSES:
            raise PlatformError(409, "RUN_NOT_DELETABLE")
        self.database.execute("BEGIN IMMEDIATE")
        try:
            self.database.execute("DELETE FROM deliveries WHERE run_id = ?", (run_id,))
            self.database.execute("DELETE FROM flow_outputs WHERE run_id = ?", (run_id,))
            self.database.execute("DELETE FROM platform_artifacts WHERE run_id = ?", (run_id,))
            self.database.execute("DELETE FROM platform_run_events WHERE run_id = ?", (run_id,))
            self.database.execute("DELETE FROM platform_runs WHERE id = ? AND project_id = ?", (run_id, project_id))
            self.database.execute("COMMIT")
        except Exception:
            self.database.execute("ROLLBACK")
            raise
        return {"runId": run_id, "deleted": True}

    def delete_runs(self, project_id: str, run_ids: list[str]) -> dict[str, Any]:
        if not run_ids:
            return {"runIds": [], "deletedCount": 0}
        placeholders = ",".join("?" for _ in run_ids)
        terminal_placeholders = ",".join("?" for _ in _TERMINAL_RUN_STATUSES)
        rows = self.database.execute(
            f"""
            SELECT id FROM platform_runs
            WHERE project_id = ?
              AND id IN ({placeholders})
              AND status IN ({terminal_placeholders})
            """,
            [project_id, *run_ids, *_TERMINAL_RUN_STATUSES],
        ).fetchall()
        deletable_run_ids = [row[0] for row in rows]
        if not deletable_run_ids:
            return {"runIds": [], "deletedCount": 0}
        deletable_placeholders = ",".join("?" for _ in deletable_run_ids)
        self.database.execute("BEGIN IMMEDIATE")
        try:
            self.database.execute(
                f"DELETE FROM deliveries WHERE run_id IN ({deletable_placeholders})",
                deletable_run_ids,
            )
            self.database.execute(
                f"DELETE FROM flow_outputs WHERE run_id IN ({deletable_placeholders})",
                deletable_run_ids,
            )
            self.database.execute(
                f"DELETE FROM platform_artifacts WHERE run_id IN ({deletable_placeholders})",
                deletable_run_ids,
            )
            self.database.execute(
                f"DELETE FROM platform_run_events WHERE run_id IN ({deletable_placeholders})",
                deletable_run_ids,
            )
            cursor = self.database.execute(
                f"DELETE FROM platform_runs WHERE id IN ({deletable_placeholders}) AND project_id = ?",
                [*deletable_run_ids, project_id],
            )
            deleted_count = cursor.rowcount
            self.database.execute("COMMIT")
        except Exception:
            self.database.execute("ROLLBACK")
            raise
        return {"runIds": deletable_run_ids, "deletedCount": deleted_count}

    def run_response(self, run: dict[str, Any]) -> dict[str, Any]:
        agent = self.database.execute(
            """
            SELECT id, name, browser_version, os, max_concurrency, last_seen_at
            FROM agents WHERE id = ?
            """,
            (run["agentId"],),
        ).fetchone()
        artifacts = self.database.execute(
            """
            SELECT id, name, content_type, created_at FROM platform_artifacts
            WHERE run_id = ? ORDER BY created_at ASC
            """,
            (run["id"],),
        ).fetchall()
        events = self.database.execute(
            """
            SELECT id, kind, data, created_at FROM platform_run_events
            WHERE run_id = ? ORDER BY id ASC LIMIT 500
            """,
            (run["id"],),
        ).fetchall()
        flow_outputs = self.database.execute(
            """
            SELECT name, value, source, created_at FROM flow_outputs
            WHERE run_id = ? ORDER BY name ASC
            """,
            (run["id"],),
        ).fetchall()
        response: dict[str, Any] = {
            **run,
            "artifacts": [
                {
                    "id": row[0],
                    "name": row[1],
                    "contentType": row[2],
                    "createdAt": row[3],
                }
                for row in artifacts
            ],
            "events": [
                {
                    "id": row[0],
                    "kind": row[1],
                    "data": parse_json(row[2], {}),
                    "at": row[3],
                }
                for row in events
            ],
            "flowOutputs": [
                {
                    "name": row[0],
                    "value": row[1],
                    "source": row[2],
                    "createdAt": row[3],
                }
                for row in flow_outputs
            ],
        }
        if run["executorType"] == "agent" and agent:
            response["agent"] = {
                "id": agent[0],
                "name": agent[1],
                "browserVersion": agent[2],
                "os": agent[3],
                "maxConcurrency": agent[4],
                "lastSeenAt": agent[5],
            }
        return response

    def cancel_managed_run(self, run_id: str) -> bool:
        return self.managed_runner.cancel(run_id)

    def managed_runner_input(self, run: dict[str, Any]) -> dict[str, Any]:
        snapshot = run["snapshot"]
        flow = as_record(snapshot.get("flow"))
        environment = as_record(snapshot.get("environment"))
        variable_rows = self.database.execute(
            """
            SELECT data FROM project_resources
            WHERE project_id = ? AND resource_type = 'variables'
              AND archived_at IS NULL
            """,
            (run["projectId"],),
        ).fetchall()
        variables: dict[str, str] = {}
        for variable_row in variable_rows:
            variable = parse_json(variable_row[0], {})
            if (
                not isinstance(variable, dict)
                or variable.get("secret") is True
                or not isinstance(variable.get("name"), str)
                or not isinstance(variable.get("value"), str)
            ):
                continue
            scope = (
                "env"
                if variable.get("scope") == "环境"
                else "project"
                if variable.get("scope") == "项目"
                else ""
            )
            key = f"{scope}.{variable['name']}" if scope else variable["name"]
            variables[key] = variable["value"]
        secret_names = snapshot.get("secretNames", [])
        if not isinstance(secret_names, list):
            secret_names = []
        secret_names = [name for name in secret_names if isinstance(name, str)]
        dataset_row = as_record(as_record(snapshot.get("datasetRow")).get("data"))
        return {
            "environment": environment,
            "flow": {
                "id": flow.get("id") if isinstance(flow.get("id"), str) else run["revisionId"],
                "name": flow.get("name") if isinstance(flow.get("name"), str) else "Published flow",
                "steps": flow.get("steps") if isinstance(flow.get("steps"), list) else [],
            },
            "elements": snapshot.get("elements")
            if isinstance(snapshot.get("elements"), list)
            else [],
            "variables": variables,
            "data": {
                str(key): str(value or "")
                for key, value in dataset_row.items()
            },
            "secrets": self.secret_values(run["projectId"], secret_names),
            "upToStepId": (
                snapshot.get("upToStepId")
                if isinstance(snapshot.get("upToStepId"), str)
                else None
            ),
        }

    def enqueue_managed_run(self, run_id: str) -> None:
        run = self.run_by_id(run_id)
        if run["status"] != "queued":
            return
        input = self.managed_runner_input(run)

        def started() -> None:
            self.database.execute(
                """
                UPDATE platform_runs SET status = 'running', updated_at = ?
                WHERE id = ? AND status = 'queued'
                """,
                (now(), run_id),
            )
            self.append_run_event(run_id, "run.started", {"executorType": "managed"})

        def event(kind: str, data: dict[str, Any]) -> None:
            self.append_run_event(
                run_id,
                kind,
                self.redact_run_value(run, data),
            )

        def artifact(input_data: dict[str, Any]) -> None:
            artifact_id = str(uuid.uuid4())
            self.database.execute(
                """
                INSERT INTO platform_artifacts (
                  id, run_id, project_id, name, content_type, path, created_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    artifact_id,
                    run_id,
                    run["projectId"],
                    safe_artifact_name(str(input_data["name"])),
                    str(input_data["contentType"])[:120],
                    str(input_data["path"]),
                    now(),
                ),
            )
            self.append_run_event(
                run_id,
                "artifact.created",
                {
                    "artifactId": artifact_id,
                    "name": str(input_data["name"]),
                    "contentType": str(input_data["contentType"]),
                },
            )

        def completed(result: dict[str, Any]) -> None:
            current_run = self.run_by_id(run_id)
            safe_result = self.redact_run_value(current_run, result)
            requested_status = (
                result.get("status")
                if result.get("status") in ("success", "failed")
                else "failed"
            )
            status = "canceled" if current_run["cancellationRequested"] else requested_status
            updated = self.database.execute(
                """
                UPDATE platform_runs
                SET status = ?, result = ?, updated_at = ?
                WHERE id = ? AND status IN ('queued', 'running')
                """,
                (status, json(safe_result), now(), run_id),
            )
            if updated.rowcount != 1:
                return
            self.persist_flow_outputs(current_run, safe_result)
            self.append_run_event(
                run_id,
                "run.complete",
                {"status": status, "result": safe_result, "executorType": "managed"},
            )
            self.audit_run_lifecycle(run_id, current_run, status)
            self.queue_run_deliveries(self.run_by_id(run_id), status)

        self.managed_runner.enqueue(
            run_id,
            input,
            {
                "started": started,
                "event": event,
                "artifact": artifact,
                "completed": completed,
            },
            kind="run",
            workspace_id=self.project_for(run["projectId"])["workspace_id"],
        )

    def metrics(self) -> dict[str, Any]:
        """OBS-02 service-level metrics (DB-backed, JSON)."""
        run_counts: dict[str, int] = {}
        for row in self.database.execute(
            "SELECT status, COUNT(*) FROM platform_runs GROUP BY status"
        ).fetchall():
            run_counts[str(row[0])] = int(row[1])

        delivery_counts: dict[str, int] = {}
        for row in self.database.execute(
            "SELECT status, COUNT(*) FROM deliveries GROUP BY status"
        ).fetchall():
            delivery_counts[str(row[0])] = int(row[1])

        disk: dict[str, int] | None = None
        try:
            usage = shutil.disk_usage(str(self.data_directory))
            disk = {"total": usage.total, "used": usage.used, "free": usage.free}
        except Exception:
            pass

        return {
            "runs": run_counts,
            "deliveries": delivery_counts,
            "disk": disk,
            "artifactBytes": self._artifact_bytes(),
        }

    def _artifact_bytes(self) -> int:
        total = 0
        try:
            for path in self.managed_runner.artifact_directory.rglob("*"):
                if path.is_file():
                    total += path.stat().st_size
        except Exception:
            pass
        return total

    def redact_run_value(self, run: dict[str, Any], value: Any) -> Any:
        try:
            rows = self.database.execute(
                """
                SELECT name, iv, tag, ciphertext FROM project_secrets
                WHERE project_id = ?
                """,
                (run["projectId"],),
            ).fetchall()
            secrets = {
                row[0]: self.decrypt({"iv": row[1], "tag": row[2], "ciphertext": row[3]})
                for row in rows
            }

            def redact(current: Any) -> Any:
                if isinstance(current, str):
                    result = current
                    for secret in secrets.values():
                        if secret:
                            result = result.replace(secret, "***")
                    return result
                if isinstance(current, list):
                    return [redact(item) for item in current]
                if isinstance(current, dict):
                    return {key: redact(item) for key, item in current.items()}
                return current

            return redact(value)
        except Exception:
            return "***"

    def build_assertion_report(
        self, run_id: str, run_format: str
    ) -> dict[str, Any]:
        """装配断言报告并登记为 run 的 artifact（JSON/XLSX）。

        数据源：run.result.assertions + 同 run 的失败截图/trace artifact 引用
        （截图名 `failure-step-{序号}.png`、trace 名 `trace.zip`，缺失留空不报错）。
        actual 经 redact_run_value 脱敏。无断言 raise 409。
        """
        from ..http import PlatformError

        run = self.run_by_id(run_id)
        result = as_record(run.get("result"))
        assertions = result.get("assertions") if isinstance(result, dict) else None
        if not isinstance(assertions, list) or not assertions:
            raise PlatformError(409, "RUN_HAS_NO_ASSERTIONS")

        artifact_rows = self.database.execute(
            """
            SELECT id, name, content_type, path FROM platform_artifacts
            WHERE run_id = ? ORDER BY created_at ASC
            """,
            (run_id,),
        ).fetchall()
        screenshots: dict[str, str] = {}
        trace_id: str | None = None
        for artifact_id, name, _content_type, _path in artifact_rows:
            if name.startswith("failure-step-") and name.endswith(".png"):
                stem = name[len("failure-step-"):][:-len(".png")]
                if stem.isdigit():
                    screenshots[stem] = artifact_id
            elif name == "trace.zip":
                trace_id = artifact_id

        snapshot = as_record(run.get("snapshot"))
        flow = as_record(snapshot.get("flow"))
        environment = as_record(snapshot.get("environment"))
        rows: list[dict[str, Any]] = []
        for item in assertions:
            if not isinstance(item, dict) or not isinstance(item.get("passed"), bool):
                continue
            step_index = int(item.get("stepIndex") or 0)
            actual = self.redact_run_value(run, item.get("actual"))
            rows.append(
                {
                    "stepIndex": step_index,
                    "stepId": str(item.get("stepId") or ""),
                    "title": str(item.get("title") or "断言"),
                    "type": str(item.get("type") or ""),
                    "passed": item["passed"],
                    "expected": str(item.get("expected") or ""),
                    "actual": str(actual) if actual is not None else "",
                    "durationMs": int(item.get("durationMs") or 0),
                    "screenshotArtifactId": screenshots.get(str(step_index + 1)),
                    "traceArtifactId": trace_id,
                }
            )
        report = {
            "runId": run_id,
            "flowName": str(flow.get("name") or "Published flow"),
            "environmentName": str(environment.get("name") or ""),
            "status": run["status"],
            "generatedAt": now(),
            "assertionCount": len(rows),
            "assertions": rows,
        }

        if run_format == "xlsx":
            content_type = (
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
            extension = "xlsx"
            data = self._assertion_report_xlsx(report)
        else:
            content_type = "application/json"
            extension = "json"
            data = _json.dumps(report, ensure_ascii=False, indent=2).encode("utf-8")

        artifact_directory = self.managed_runner.artifact_directory
        artifact_directory.mkdir(parents=True, exist_ok=True)
        artifact_name = safe_artifact_name(f"assertion-report-{run_id}.{extension}")
        artifact_path = artifact_directory / artifact_name
        artifact_path.write_bytes(data)
        artifact_id = str(uuid.uuid4())
        created_at = now()
        self.database.execute(
            """
            INSERT INTO platform_artifacts (
              id, run_id, project_id, name, content_type, path, created_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?)
            """,
            (
                artifact_id,
                run_id,
                run["projectId"],
                artifact_name,
                content_type,
                str(artifact_path),
                created_at,
            ),
        )
        return {
            "artifact": {
                "id": artifact_id,
                "name": artifact_name,
                "contentType": content_type,
                "createdAt": created_at,
            }
        }

    def _assertion_report_xlsx(self, report: dict[str, Any]) -> bytes:
        import io

        from openpyxl import Workbook
        from openpyxl.styles import Font

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "断言报告"
        sheet.append(["序号", "步骤", "类型", "判定", "期望", "实际", "耗时(ms)", "失败截图", "Trace"])
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        for row in report["assertions"]:
            sheet.append(
                [
                    int(row["stepIndex"]) + 1,
                    row["title"],
                    row["type"],
                    "通过" if row["passed"] else "失败",
                    row["expected"],
                    row["actual"],
                    int(row["durationMs"]),
                    row["screenshotArtifactId"] or "",
                    row["traceArtifactId"] or "",
                ]
            )
        buffer = io.BytesIO()
        workbook.save(buffer)
        return buffer.getvalue()

    def assertion_stats_for_runs(self, runs: list[dict[str, Any]]) -> dict[str, int]:
        """跨 run 应用层聚合断言计数（口径写死：分子=含断言 run 中 passed 总数，
        分母=含断言 run 的断言总数；无断言 run 不进分子分母）。"""
        runs_with_assertions = 0
        total = 0
        passed = 0
        for run in runs:
            result = as_record(run.get("result"))
            assertions = result.get("assertions") if isinstance(result, dict) else None
            if not isinstance(assertions, list):
                continue
            total_this = 0
            passed_this = 0
            for item in assertions:
                if isinstance(item, dict) and isinstance(item.get("passed"), bool):
                    total_this += 1
                    if item["passed"]:
                        passed_this += 1
            if total_this > 0:
                runs_with_assertions += 1
                total += total_this
                passed += passed_this
        return {
            "runsWithAssertions": runs_with_assertions,
            "totalAssertions": total,
            "passedAssertions": passed,
            "failedAssertions": total - passed,
        }

    def assertion_failures_for_runs(
        self, runs: list[dict[str, Any]]
    ) -> list[dict[str, Any]]:
        """跨 run 收集失败断言明细（batch detail 的「失败明细列表」数据源）。

        actual 经 redact_run_value 脱敏，与断言载荷脱敏约束一致。
        """
        failures: list[dict[str, Any]] = []
        for run in runs:
            result = as_record(run.get("result"))
            assertions = result.get("assertions") if isinstance(result, dict) else None
            if not isinstance(assertions, list):
                continue
            snapshot = as_record(run.get("snapshot"))
            flow = as_record(snapshot.get("flow"))
            flow_name = str(flow.get("name") or "Published flow")
            for item in assertions:
                if not isinstance(item, dict) or item.get("passed") is not False:
                    continue
                actual = self.redact_run_value(run, item.get("actual"))
                failures.append(
                    {
                        "runId": str(run.get("id") or ""),
                        "flowName": flow_name,
                        "title": str(item.get("title") or "断言"),
                        "type": str(item.get("type") or ""),
                        "expected": str(item.get("expected") or ""),
                        "actual": str(actual) if actual is not None else "",
                    }
                )
        return failures

    def assertion_stats(
        self, project_id: str, window_days: int | None = None
    ) -> dict[str, Any]:
        """项目级断言聚合（全量扫描含断言 run，非分页口径）。

        SQLite 对 JSON 列聚合不友好，应用层扫描 `platform_runs.result` 累加。
        window_days 为 None 或 <=0 时不做时间过滤（全量）。返回含 windowDays。
        """
        params: list[Any] = [project_id]
        window_sql = ""
        if window_days is not None and window_days > 0:
            window_sql = "AND created_at >= ?"
            params.append(days_ago_iso(window_days))
        rows = self.database.execute(
            f"""
            SELECT result FROM platform_runs
            WHERE project_id = ? {window_sql}
            """,
            tuple(params),
        ).fetchall()
        runs = [{"result": parse_json(row[0], None)} for row in rows]
        stats = self.assertion_stats_for_runs(runs)
        stats["windowDays"] = window_days
        return stats

    def persist_flow_outputs(
        self,
        run: dict[str, Any],
        result: dict[str, Any],
    ) -> None:
        outputs = as_record(result.get("flowOutputs"))
        allowed_names = public_flow_output_names(run)
        for name, source_value in outputs.items():
            output_name = str(name).strip()[:120]
            if output_name not in allowed_names:
                continue
            if isinstance(source_value, str):
                value = source_value
            elif isinstance(source_value, (int, float, bool)):
                value = str(source_value)
            else:
                value = ""
            value = self.redact_run_value(run, value)
            if not value:
                continue
            self.database.execute(
                """
                INSERT INTO flow_outputs (
                  id, run_id, name, value, source, created_at
                ) VALUES (?, ?, ?, ?, 'agent', ?)
                ON CONFLICT(run_id, name) DO UPDATE SET
                  value = excluded.value, source = excluded.source,
                  created_at = excluded.created_at
                """,
                (str(uuid.uuid4()), run["id"], output_name, str(value)[:20000], now()),
            )

    def audit_run_lifecycle(
        self,
        run_id: str,
        run: dict[str, Any],
        status: str,
    ) -> None:
        project = self.project_for(run["projectId"])
        detail: dict[str, Any] = {"status": status}
        if status == "failed":
            failure = self.database.execute(
                """
                SELECT kind, data FROM platform_run_events
                WHERE run_id = ? AND (kind LIKE '%failed%' OR kind LIKE '%error%')
                ORDER BY id DESC LIMIT 1
                """,
                (run_id,),
            ).fetchone()
            failure_data = parse_json(failure[1], {}) if failure else {}
            if isinstance(failure_data, dict):
                if isinstance(failure_data.get("code"), str) and failure_data["code"]:
                    detail["errorCode"] = failure_data["code"]
                elif isinstance(failure_data.get("reason"), str) and failure_data["reason"]:
                    detail["errorCode"] = failure_data["reason"][:200]
                if isinstance(failure_data.get("stepId"), str) and failure_data["stepId"]:
                    detail["stepId"] = failure_data["stepId"]
        self.audit(
            project["workspace_id"],
            {"type": "system", "id": "managed-runner"},
            f"run.{'canceled' if status == 'canceled' else 'failed' if status == 'failed' else 'completed'}",
            {"type": "run_batch", "id": run_id},
            detail,
            run["projectId"],
        )

    def finalize_run_as_interrupted(self, run_id: str, reason: str) -> None:
        updated = self.database.execute(
            """
            UPDATE platform_runs
            SET status = 'failed', result = ?, updated_at = ?
            WHERE id = ? AND status IN ('queued', 'running')
            """,
            (json({"error": reason, "interrupted": True}), now(), run_id),
        )
        if updated.rowcount != 1:
            return
        self.append_run_event(run_id, "run.interrupted", {"reason": reason})
        self.append_run_event(
            run_id, "run.failed", {"reason": reason, "interrupted": True}
        )
        self.audit_run_lifecycle(run_id, self.run_by_id(run_id), "failed")
        self.queue_run_deliveries(self.run_by_id(run_id), "failed")
