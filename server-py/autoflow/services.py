"""Platform service layer matching server/platform.ts."""

from __future__ import annotations

import json as _json
import os
import sqlite3
import time as _time
import uuid
import base64
import concurrent.futures
import http.client
import io
import ipaddress
import re
import socket
import shutil
import ssl
from urllib.parse import urljoin, urlsplit
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

from .audit import create_audit_writer, create_deployment_audit_writer
from .core import (
    digest,
    failure_category,
    json,
    next_cron_time,
    normalize_dataset_rows,
    notification_rejection_code,
    now,
    parse_csv,
    parse_json,
    public_flow_output_names,
    safe_artifact_name,
)
from .crypto import decrypt, encrypt, key_material
from .migrations import migrate_project_document_resources, run_platform_migrations
from .managed_runner import ManagedRunner
from .recorder import RecordingCoordinator
from .recording_state import RecordingSessionStateStore
from .resources import as_record
from .workspaces import (
    GLOBAL_ROLE_SUPER_ADMIN,
    WORKSPACE_ROLE_ADMIN,
    capabilities_for_role,
    is_super_admin,
    is_workspace_role,
    normalize_workspace_role,
    role_has_capability,
)


_CHROMIUM_AVAILABLE: bool | None = None
_CHROMIUM_AVAILABLE_AT: float | None = None
_CHROMIUM_CACHE_TTL_NEGATIVE: float = 30.0  # 不可用结果 30s 后允许重试，避免安装完 Chromium 还要重启进程
_CHROMIUM_CACHE_TTL_POSITIVE: float = 3600.0  # 可用结果缓存 1h
_TERMINAL_RUN_STATUSES = ("success", "failed", "canceled")


BOOTSTRAP_SCHEMA = """
    CREATE TABLE IF NOT EXISTS platform_users (
      id TEXT PRIMARY KEY,
      email TEXT NOT NULL UNIQUE,
      name TEXT NOT NULL,
      created_at TEXT NOT NULL
    );
    CREATE TABLE IF NOT EXISTS platform_user_credentials (
      user_id TEXT PRIMARY KEY REFERENCES platform_users(id),
      password_hash TEXT NOT NULL,
      created_at TEXT NOT NULL,
      updated_at TEXT NOT NULL
    );
    CREATE TABLE IF NOT EXISTS platform_sessions (
      token_hash TEXT PRIMARY KEY,
      user_id TEXT NOT NULL REFERENCES platform_users(id),
      expires_at TEXT NOT NULL,
      created_at TEXT NOT NULL
    );
    CREATE TABLE IF NOT EXISTS workspaces (
      id TEXT PRIMARY KEY,
      name TEXT NOT NULL,
      created_at TEXT NOT NULL
    );
    CREATE TABLE IF NOT EXISTS workspace_members (
      workspace_id TEXT NOT NULL REFERENCES workspaces(id),
      user_id TEXT NOT NULL REFERENCES platform_users(id),
      role TEXT NOT NULL,
      PRIMARY KEY (workspace_id, user_id)
    );
    CREATE TABLE IF NOT EXISTS platform_projects (
      id TEXT PRIMARY KEY,
      workspace_id TEXT NOT NULL REFERENCES workspaces(id),
      slug TEXT NOT NULL,
      name TEXT NOT NULL,
      description TEXT NOT NULL,
      archived_at TEXT,
      created_at TEXT NOT NULL,
      updated_at TEXT NOT NULL,
      UNIQUE (workspace_id, slug)
    );
    CREATE TABLE IF NOT EXISTS project_documents (
      project_id TEXT PRIMARY KEY REFERENCES platform_projects(id),
      data TEXT NOT NULL,
      version INTEGER NOT NULL,
      updated_at TEXT NOT NULL
    );
    CREATE TABLE IF NOT EXISTS platform_imports (
      id TEXT PRIMARY KEY,
      workspace_id TEXT NOT NULL REFERENCES workspaces(id),
      source_id TEXT NOT NULL,
      imported_at TEXT NOT NULL,
      result TEXT NOT NULL,
      UNIQUE (workspace_id, source_id)
    );
    CREATE TABLE IF NOT EXISTS flow_revisions (
      id TEXT PRIMARY KEY,
      project_id TEXT NOT NULL REFERENCES platform_projects(id),
      flow_id TEXT,
      flow_name TEXT,
      environment_id TEXT,
      revision_number INTEGER NOT NULL,
      status TEXT NOT NULL,
      flow_snapshot TEXT NOT NULL,
      environment_snapshot TEXT NOT NULL,
      element_snapshot TEXT NOT NULL,
      dataset_snapshot TEXT NOT NULL,
      checksum TEXT NOT NULL,
      created_by TEXT NOT NULL,
      created_at TEXT NOT NULL,
      published_at TEXT,
      UNIQUE (project_id, revision_number)
    );
    CREATE TABLE IF NOT EXISTS project_secrets (
      id TEXT PRIMARY KEY,
      project_id TEXT NOT NULL REFERENCES platform_projects(id),
      name TEXT NOT NULL,
      key_version INTEGER NOT NULL,
      iv TEXT NOT NULL,
      tag TEXT NOT NULL,
      ciphertext TEXT NOT NULL,
      created_at TEXT NOT NULL,
      updated_at TEXT NOT NULL,
      UNIQUE (project_id, name)
    );
    CREATE TABLE IF NOT EXISTS audit_events (
      id TEXT PRIMARY KEY,
      workspace_id TEXT NOT NULL REFERENCES workspaces(id),
      project_id TEXT,
      actor_type TEXT NOT NULL,
      actor_id TEXT NOT NULL,
      action TEXT NOT NULL,
      target_type TEXT NOT NULL,
      target_id TEXT NOT NULL,
      detail TEXT NOT NULL,
      created_at TEXT NOT NULL
    );
    CREATE TABLE IF NOT EXISTS deployment_audit_events (
      id TEXT PRIMARY KEY,
      actor_type TEXT NOT NULL,
      actor_id TEXT NOT NULL,
      action TEXT NOT NULL,
      target_type TEXT NOT NULL,
      target_id TEXT NOT NULL,
      detail TEXT NOT NULL,
      created_at TEXT NOT NULL
    );
    CREATE INDEX IF NOT EXISTS deployment_audit_events_created
      ON deployment_audit_events (created_at DESC);
    CREATE TABLE IF NOT EXISTS agents (
      id TEXT PRIMARY KEY,
      workspace_id TEXT NOT NULL REFERENCES workspaces(id),
      name TEXT NOT NULL,
      credential_hash TEXT NOT NULL,
      status TEXT NOT NULL,
      browser_version TEXT NOT NULL,
      os TEXT NOT NULL,
      max_concurrency INTEGER NOT NULL,
      current_task TEXT,
      last_seen_at TEXT,
      created_at TEXT NOT NULL,
      revoked_at TEXT
    );
    CREATE TABLE IF NOT EXISTS platform_runs (
      id TEXT PRIMARY KEY,
      project_id TEXT NOT NULL REFERENCES platform_projects(id),
      revision_id TEXT NOT NULL REFERENCES flow_revisions(id),
      environment_id TEXT NOT NULL,
      agent_id TEXT NOT NULL REFERENCES agents(id),
      status TEXT NOT NULL,
      snapshot TEXT NOT NULL,
      cancellation_requested INTEGER NOT NULL DEFAULT 0,
      result TEXT,
      created_by TEXT NOT NULL,
      created_at TEXT NOT NULL,
      updated_at TEXT NOT NULL
    );
    CREATE TABLE IF NOT EXISTS element_validations (
      id TEXT PRIMARY KEY,
      project_id TEXT NOT NULL REFERENCES platform_projects(id),
      environment_id TEXT NOT NULL,
      agent_id TEXT NOT NULL REFERENCES agents(id),
      status TEXT NOT NULL,
      element_snapshot TEXT NOT NULL,
      result TEXT,
      error TEXT,
      created_by TEXT NOT NULL,
      created_at TEXT NOT NULL,
      updated_at TEXT NOT NULL
    );
    CREATE TABLE IF NOT EXISTS platform_run_events (
      id INTEGER PRIMARY KEY AUTOINCREMENT,
      run_id TEXT NOT NULL REFERENCES platform_runs(id),
      kind TEXT NOT NULL,
      data TEXT NOT NULL,
      created_at TEXT NOT NULL
    );
    CREATE TABLE IF NOT EXISTS platform_artifacts (
      id TEXT PRIMARY KEY,
      run_id TEXT NOT NULL REFERENCES platform_runs(id),
      project_id TEXT NOT NULL REFERENCES platform_projects(id),
      name TEXT NOT NULL,
      content_type TEXT NOT NULL,
      path TEXT NOT NULL,
      created_at TEXT NOT NULL
    );
    CREATE TABLE IF NOT EXISTS datasets (
      id TEXT PRIMARY KEY,
      project_id TEXT NOT NULL REFERENCES platform_projects(id),
      name TEXT NOT NULL,
      description TEXT NOT NULL,
      created_by TEXT NOT NULL,
      created_at TEXT NOT NULL,
      updated_at TEXT NOT NULL,
      UNIQUE (project_id, name)
    );
    CREATE TABLE IF NOT EXISTS dataset_versions (
      id TEXT PRIMARY KEY,
      dataset_id TEXT NOT NULL REFERENCES datasets(id),
      version_number INTEGER NOT NULL,
      columns_json TEXT NOT NULL,
      row_count INTEGER NOT NULL,
      checksum TEXT NOT NULL,
      source_name TEXT NOT NULL,
      created_by TEXT NOT NULL,
      created_at TEXT NOT NULL,
      UNIQUE (dataset_id, version_number)
    );
    CREATE TABLE IF NOT EXISTS dataset_rows (
      id TEXT PRIMARY KEY,
      dataset_version_id TEXT NOT NULL REFERENCES dataset_versions(id),
      row_number INTEGER NOT NULL,
      data_json TEXT NOT NULL,
      UNIQUE (dataset_version_id, row_number)
    );
    CREATE TABLE IF NOT EXISTS flow_outputs (
      id TEXT PRIMARY KEY,
      run_id TEXT NOT NULL REFERENCES platform_runs(id),
      name TEXT NOT NULL,
      value TEXT NOT NULL,
      source TEXT NOT NULL,
      created_at TEXT NOT NULL,
      UNIQUE (run_id, name)
    );
    CREATE TABLE IF NOT EXISTS schedules (
      id TEXT PRIMARY KEY,
      project_id TEXT NOT NULL REFERENCES platform_projects(id),
      revision_id TEXT NOT NULL REFERENCES flow_revisions(id),
      environment_id TEXT NOT NULL,
      dataset_version_id TEXT REFERENCES dataset_versions(id),
      name TEXT NOT NULL,
      cron_expression TEXT NOT NULL,
      timezone TEXT NOT NULL,
      enabled INTEGER NOT NULL DEFAULT 1,
      last_run_at TEXT,
      next_run_at TEXT NOT NULL,
      created_by TEXT NOT NULL,
      created_at TEXT NOT NULL,
      updated_at TEXT NOT NULL
    );
    CREATE TABLE IF NOT EXISTS webhook_triggers (
      id TEXT PRIMARY KEY,
      project_id TEXT NOT NULL REFERENCES platform_projects(id),
      revision_id TEXT NOT NULL REFERENCES flow_revisions(id),
      environment_id TEXT NOT NULL,
      dataset_version_id TEXT REFERENCES dataset_versions(id),
      name TEXT NOT NULL,
      signing_secret_iv TEXT,
      signing_secret_tag TEXT,
      signing_secret_ciphertext TEXT,
      enabled INTEGER NOT NULL DEFAULT 1,
      created_by TEXT NOT NULL,
      created_at TEXT NOT NULL,
      last_triggered_at TEXT
    );
    CREATE TABLE IF NOT EXISTS webhook_deliveries (
      trigger_id TEXT NOT NULL REFERENCES webhook_triggers(id),
      delivery_id TEXT NOT NULL,
      received_at TEXT NOT NULL,
      PRIMARY KEY (trigger_id, delivery_id)
    );
    CREATE TABLE IF NOT EXISTS notification_channels (
      id TEXT PRIMARY KEY,
      workspace_id TEXT NOT NULL REFERENCES workspaces(id),
      name TEXT NOT NULL,
      channel_type TEXT NOT NULL,
      config_iv TEXT NOT NULL,
      config_tag TEXT NOT NULL,
      config_ciphertext TEXT NOT NULL,
      enabled INTEGER NOT NULL DEFAULT 1,
      created_by TEXT NOT NULL,
      created_at TEXT NOT NULL,
      updated_at TEXT NOT NULL,
      UNIQUE (workspace_id, name)
    );
    CREATE TABLE IF NOT EXISTS notification_subscriptions (
      project_id TEXT NOT NULL REFERENCES platform_projects(id),
      channel_id TEXT NOT NULL REFERENCES notification_channels(id),
      on_success INTEGER NOT NULL DEFAULT 0,
      on_failure INTEGER NOT NULL DEFAULT 1,
      PRIMARY KEY (project_id, channel_id)
    );
    CREATE TABLE IF NOT EXISTS deliveries (
      id TEXT PRIMARY KEY,
      channel_id TEXT NOT NULL REFERENCES notification_channels(id),
      run_id TEXT NOT NULL REFERENCES platform_runs(id),
      status TEXT NOT NULL,
      attempt_count INTEGER NOT NULL DEFAULT 0,
      payload TEXT NOT NULL,
      response_code INTEGER,
      error TEXT,
      next_attempt_at TEXT,
      created_at TEXT NOT NULL,
      delivered_at TEXT,
      updated_at TEXT NOT NULL,
      UNIQUE (channel_id, run_id)
    );
    CREATE INDEX IF NOT EXISTS platform_projects_workspace ON platform_projects (workspace_id, archived_at);
    CREATE INDEX IF NOT EXISTS flow_revisions_project ON flow_revisions (project_id, revision_number DESC);
    CREATE INDEX IF NOT EXISTS agents_workspace ON agents (workspace_id, status, last_seen_at);
    CREATE INDEX IF NOT EXISTS platform_runs_project ON platform_runs (project_id, created_at DESC);
    CREATE INDEX IF NOT EXISTS datasets_project ON datasets (project_id, updated_at DESC);
    CREATE INDEX IF NOT EXISTS dataset_versions_dataset ON dataset_versions (dataset_id, version_number DESC);
    CREATE INDEX IF NOT EXISTS dataset_rows_version ON dataset_rows (dataset_version_id, row_number);
    CREATE INDEX IF NOT EXISTS flow_outputs_run ON flow_outputs (run_id, name);
    CREATE INDEX IF NOT EXISTS schedules_due ON schedules (enabled, next_run_at);
    CREATE INDEX IF NOT EXISTS webhook_triggers_project ON webhook_triggers (project_id, enabled);
    CREATE INDEX IF NOT EXISTS webhook_deliveries_received ON webhook_deliveries (received_at);
    CREATE INDEX IF NOT EXISTS deliveries_channel ON deliveries (channel_id, status, created_at DESC);
    CREATE INDEX IF NOT EXISTS audit_events_project ON audit_events (project_id, created_at DESC);
    CREATE INDEX IF NOT EXISTS platform_run_events_run ON platform_run_events (run_id, id);
    CREATE INDEX IF NOT EXISTS deliveries_due ON deliveries (status, next_attempt_at);
"""


@dataclass
class AuthUser:
    id: str
    email: str
    name: str
    global_role: str | None = None


def _format_notification_body(
    channel_type: str,
    payload: dict[str, Any],
    keyword: str | None = None,
) -> dict[str, Any]:
    content = (
        f"{keyword} AutoFlow {payload.get('status', '')}: "
        f"{payload.get('runId', '')} ({payload.get('environmentId', '')})"
        if keyword
        else (
            f"AutoFlow {payload.get('status', '')}: "
            f"{payload.get('runId', '')} ({payload.get('environmentId', '')})"
        )
    )
    if channel_type == "feishu":
        return {"msg_type": "text", "content": {"text": content}}
    if channel_type == "dingtalk":
        return {"msgtype": "text", "text": {"content": content}}
    if channel_type == "wecom":
        return {"msgtype": "text", "text": {"content": content}}
    return payload


def _post_notification(
    target: dict[str, str],
    headers: dict[str, str],
    body: str,
) -> dict[str, Any]:
    parsed = urlsplit(target["url"])
    hostname = parsed.hostname or ""
    port = parsed.port or (443 if parsed.scheme == "https" else 80)
    path = parsed.path or "/"
    if parsed.query:
        path = f"{path}?{parsed.query}"
    request_headers = {
        "Host": parsed.netloc or hostname,
        "content-type": "application/json",
        **headers,
    }
    if parsed.scheme == "https":
        connection = http.client.HTTPSConnection(
            target["address"],
            port,
            timeout=10,
            context=ssl.create_default_context(),
            server_hostname=hostname,
        )
    else:
        connection = http.client.HTTPConnection(
            target["address"], port, timeout=10
        )
    try:
        connection.request("POST", path, body=body, headers=request_headers)
        response = connection.getresponse()
        response_body = response.read(2048).decode("utf-8", errors="replace")
        return {"status": response.status, "body": response_body[:2048]}
    finally:
        connection.close()


def _iso_add_seconds(seconds: int) -> str:
    return (
        datetime.now(timezone.utc) + timedelta(seconds=seconds)
    ).strftime("%Y-%m-%dT%H:%M:%S.%f")[:-3] + "Z"


class PlatformServices:
    def __init__(self, data_directory: str):
        self.data_directory = Path(data_directory)
        self.data_directory.mkdir(parents=True, exist_ok=True)
        self.database = sqlite3.connect(
            self.data_directory / "platform.sqlite", check_same_thread=False
        )
        self.database.isolation_level = None
        run_platform_migrations(self.database, BOOTSTRAP_SCHEMA)
        self.audit = create_audit_writer(self.database)
        self.deployment_audit = create_deployment_audit_writer(self.database)
        self.managed_runner = ManagedRunner(
            self.data_directory / "artifacts",
            global_concurrency=int(
                os.environ.get("AUTOFLOW_RUNNER_GLOBAL_CONCURRENCY", "2")
            ),
            workspace_concurrency=int(
                os.environ.get("AUTOFLOW_RUNNER_WORKSPACE_CONCURRENCY", "1")
            ),
        )
        self._recording_executor = concurrent.futures.ThreadPoolExecutor(
            max_workers=1, thread_name_prefix="recording"
        )
        self.recording_session_state = RecordingSessionStateStore()
        self.recording_coordinator = RecordingCoordinator(
            submit=self._recording_executor.submit,
            on_failed=self._audit_recording_failed,
            on_storage_state=self.recording_session_state.remember,
        )
        self.webhook_requests: dict[str, list[float]] = {}
        configured_secret = os.environ.get("PLATFORM_SECRET_KEY")
        if not configured_secret:
            key_file = os.environ.get("PLATFORM_SECRET_KEY_FILE")
            if key_file:
                try:
                    configured_secret = Path(key_file).read_text(
                        encoding="utf-8"
                    ).strip()
                except Exception:
                    configured_secret = None
        if os.environ.get("NODE_ENV") == "production" and not configured_secret:
            raise RuntimeError("PLATFORM_SECRET_KEY is required in production")
        self.key_material = key_material(configured_secret)
        self._configured_secret = configured_secret
        interrupted = self.database.execute(
            """
            SELECT id FROM platform_runs
            WHERE executor_type = 'managed' AND status = 'running'
            """
        ).fetchall()
        for row in interrupted:
            self.finalize_run_as_interrupted(row[0], "SERVICE_RESTARTED")
        recoverable = self.database.execute(
            """
            SELECT id FROM platform_runs
            WHERE executor_type = 'managed' AND status = 'queued'
            ORDER BY created_at
            """
        ).fetchall()
        for row in recoverable:
            try:
                self.enqueue_managed_run(row[0])
            except Exception:
                self.database.execute(
                    """
                    UPDATE platform_runs
                    SET status = 'failed', result = ?, updated_at = ?
                    WHERE id = ? AND status = 'queued'
                    """,
                    (json({"error": "RUN_ENQUEUE_FAILED", "interrupted": True}), now(), row[0]),
                )
                self.append_run_event(
                    row[0], "run.interrupted", {"reason": "RUN_ENQUEUE_FAILED"}
                )

    def close(self) -> None:
        try:
            self.recording_coordinator.close_all()
        except Exception:
            pass
        self._recording_executor.shutdown(wait=False)
        self.managed_runner.stop()
        self.database.close()

    def _audit_recording_failed(self, session: dict[str, Any]) -> None:
        """Record a safe lifecycle summary without exposing browser error details."""
        try:
            project = self.project_for(session["projectId"])
            self.audit(
                project["workspace_id"],
                {"type": "user", "id": session["ownerId"]},
                "recording.session_failed",
                {"type": "recording_session", "id": session["id"]},
                {
                    "flowId": session["flowId"],
                    "environmentId": session["environmentId"],
                    "status": "failed",
                },
                session["projectId"],
            )
        except Exception:
            # Browser startup failures must retain their original error contract.
            pass

    def recording_login_state(
        self, owner_id: str, project_id: str, environment_id: str
    ) -> dict[str, Any] | None:
        return self.recording_session_state.state_for(
            owner_id, project_id, environment_id
        )

    def encrypt(self, value: str) -> dict[str, str]:
        encrypted = encrypt(value, self._configured_secret)
        return {
            "iv": encrypted.iv,
            "tag": encrypted.tag,
            "ciphertext": encrypted.ciphertext,
        }

    def decrypt(self, row: dict[str, str] | Any) -> str:
        return decrypt(row, self._configured_secret)

    def allow_webhook_request(self, trigger_id: str) -> bool:
        from .core import WEBHOOK_RATE_LIMIT_PER_MINUTE

        cutoff = _now_ms() - 60_000
        requests = [
            value for value in self.webhook_requests.get(trigger_id, []) if value > cutoff
        ]
        if len(requests) >= WEBHOOK_RATE_LIMIT_PER_MINUTE:
            return False
        requests.append(_now_ms())
        self.webhook_requests[trigger_id] = requests
        return True

    def create_auth_session(self, user: AuthUser) -> dict[str, Any]:
        import secrets

        token = secrets.token_urlsafe(32)
        expires_at = _iso_add_seconds(12 * 60 * 60)
        self.database.execute(
            """
            INSERT INTO platform_sessions (token_hash, user_id, expires_at, created_at)
            VALUES (?, ?, ?, ?)
            """,
            (digest(token), user.id, expires_at, now()),
        )
        return {
            "token": token,
            "expiresAt": expires_at,
            "user": {
                "id": user.id,
                "email": user.email,
                "name": user.name,
                "globalRole": user.global_role,
            },
        }

    def session_user(self, headers: dict[str, str] | None = None) -> AuthUser:
        from .core import authorization
        from .http import PlatformError

        token = authorization(headers)
        if not token:
            raise PlatformError(401, "AUTH_REQUIRED")
        row = self.database.execute(
            """
            SELECT u.id, u.email, u.name, u.global_role
            FROM platform_sessions s
            JOIN platform_users u ON u.id = s.user_id
            WHERE s.token_hash = ? AND s.expires_at > ? AND u.enabled = 1
            """,
            (digest(token), now()),
        ).fetchone()
        if not row:
            raise PlatformError(401, "SESSION_INVALID")
        return AuthUser(row[0], row[1], row[2], row[3])

    def global_role_for_user(self, user_id: str) -> str | None:
        row = self.database.execute(
            "SELECT global_role FROM platform_users WHERE id = ? AND enabled = 1",
            (user_id,),
        ).fetchone()
        return str(row[0]) if row and row[0] else None

    def require_super_admin(self, user_id: str) -> None:
        from .http import PlatformError

        if not is_super_admin(self.global_role_for_user(user_id)):
            raise PlatformError(403, "SUPER_ADMIN_REQUIRED")

    def revoke_user_sessions(self, user_id: str) -> int:
        cursor = self.database.execute(
            "DELETE FROM platform_sessions WHERE user_id = ?", (user_id,)
        )
        return max(0, cursor.rowcount)

    def create_workspace(self, user: AuthUser, name: str) -> dict[str, Any]:
        workspace = {
            "id": str(uuid.uuid4()),
            "name": name.strip()[:120] or "My workspace",
            "createdAt": now(),
        }
        self.database.execute(
            "INSERT INTO workspaces (id, name, created_at) VALUES (?, ?, ?)",
            (workspace["id"], workspace["name"], workspace["createdAt"]),
        )
        self.database.execute(
            """
            INSERT INTO workspace_members (workspace_id, user_id, role)
            VALUES (?, ?, ?)
            """,
            (workspace["id"], user.id, WORKSPACE_ROLE_ADMIN),
        )
        self.audit(
            workspace["id"],
            {"type": "user", "id": user.id},
            "workspace.created",
            {"type": "workspace", "id": workspace["id"]},
            {"name": workspace["name"]},
        )
        return workspace

    def workspaces_for_user(self, user_id: str) -> list[dict[str, Any]]:
        global_role = self.global_role_for_user(user_id)
        if is_super_admin(global_role):
            rows = self.database.execute(
                "SELECT id, name, created_at FROM workspaces ORDER BY created_at ASC"
            ).fetchall()
            return [
                {
                    "id": row[0],
                    "name": row[1],
                    "createdAt": row[2],
                    "role": GLOBAL_ROLE_SUPER_ADMIN,
                    "capabilities": capabilities_for_role(
                        GLOBAL_ROLE_SUPER_ADMIN, global_role
                    ),
                }
                for row in rows
            ]
        rows = self.database.execute(
            """
            SELECT w.id, w.name, w.created_at, m.role
            FROM workspaces w
            JOIN workspace_members m ON m.workspace_id = w.id
            WHERE m.user_id = ?
            ORDER BY w.created_at ASC
            """,
            (user_id,),
        ).fetchall()
        return [
            {
                "id": row[0],
                "name": row[1],
                "createdAt": row[2],
                "role": normalize_workspace_role(row[3]),
                "capabilities": capabilities_for_role(str(row[3]), global_role),
            }
            for row in rows
            if is_workspace_role(normalize_workspace_role(row[3]))
        ]

    def _workspace_exists(self, workspace_id: str) -> None:
        from .http import PlatformError

        row = self.database.execute(
            "SELECT id FROM workspaces WHERE id = ?", (workspace_id,)
        ).fetchone()
        if not row:
            raise PlatformError(404, "WORKSPACE_NOT_FOUND")

    def member_role(self, workspace_id: str, user_id: str) -> str:
        from .http import PlatformError

        row = self.database.execute(
            """
            SELECT role FROM workspace_members
            WHERE workspace_id = ? AND user_id = ?
            """,
            (workspace_id, user_id),
        ).fetchone()
        role = normalize_workspace_role(str(row[0])) if row else ""
        if not is_workspace_role(role):
            raise PlatformError(403, "WORKSPACE_ACCESS_DENIED")
        return role

    def effective_workspace_role(self, workspace_id: str, user_id: str) -> str:
        self._workspace_exists(workspace_id)
        global_role = self.global_role_for_user(user_id)
        if is_super_admin(global_role):
            membership = self.database.execute(
                """
                SELECT role FROM workspace_members
                WHERE workspace_id = ? AND user_id = ?
                """,
                (workspace_id, user_id),
            ).fetchone()
            if not membership:
                self.audit(
                    workspace_id,
                    {"type": "user", "id": user_id},
                    "super_admin.workspace_accessed",
                    {"type": "workspace", "id": workspace_id},
                    {},
                )
            return GLOBAL_ROLE_SUPER_ADMIN
        return self.member_role(workspace_id, user_id)

    def require_workspace_role(
        self, workspace_id: str, user_id: str, admin: bool = False
    ) -> str:
        from .http import PlatformError

        role = self.effective_workspace_role(workspace_id, user_id)
        if admin and not role_has_capability(role, "workspace.manage"):
            raise PlatformError(403, "CAPABILITY_REQUIRED")
        return role

    def require_workspace_capability(
        self, workspace_id: str, user_id: str, capability: str
    ) -> str:
        from .http import PlatformError

        role = self.effective_workspace_role(workspace_id, user_id)
        if not role_has_capability(role, capability):
            raise PlatformError(403, "CAPABILITY_REQUIRED")
        return role

    def project_for(self, project_id: str) -> dict[str, Any]:
        from .http import PlatformError

        row = self.database.execute(
            """
            SELECT id, workspace_id, source_project_id, slug, name, description,
                   archived_at, created_at, updated_at
            FROM platform_projects WHERE id = ?
            """,
            (project_id,),
        ).fetchone()
        if not row:
            raise PlatformError(404, "PROJECT_NOT_FOUND")
        return {
            "id": row[0],
            "workspace_id": row[1],
            "source_project_id": row[2],
            "slug": row[3],
            "name": row[4],
            "description": row[5],
            "archived_at": row[6],
            "created_at": row[7],
            "updated_at": row[8],
        }

    def project_response(self, project: dict[str, Any]) -> dict[str, Any]:
        return {
            "id": project["id"],
            "workspaceId": project["workspace_id"],
            "sourceProjectId": project.get("source_project_id") or None,
            "slug": project["slug"],
            "name": project["name"],
            "description": project["description"],
            "archivedAt": project.get("archived_at"),
            "createdAt": project.get("created_at"),
            "updatedAt": project.get("updated_at"),
        }

    def require_project_role(
        self, project_id: str, user_id: str, write: bool = False
    ) -> dict[str, Any]:
        project = self.project_for(project_id)
        if write:
            role = self.require_workspace_capability(
                project["workspace_id"], user_id, "project.manage"
            )
        else:
            role = self.require_workspace_role(project["workspace_id"], user_id)
        return {"project": project, "role": role}

    def require_project_admin(self, project_id: str, user_id: str) -> dict[str, Any]:
        return self.require_project_role(project_id, user_id, True)

    def require_project_capability(
        self, project_id: str, user_id: str, capability: str
    ) -> dict[str, Any]:
        project = self.project_for(project_id)
        role = self.require_workspace_capability(
            project["workspace_id"], user_id, capability
        )
        return {"project": project, "role": role}

    def account_for(self, user_id: str) -> dict[str, Any]:
        from .http import PlatformError

        row = self.database.execute(
            """
            SELECT id, email, name, enabled, global_role, created_at
            FROM platform_users WHERE id = ?
            """,
            (user_id,),
        ).fetchone()
        if not row:
            raise PlatformError(404, "ACCOUNT_NOT_FOUND")
        return {
            "id": row[0],
            "email": row[1],
            "name": row[2],
            "enabled": bool(row[3]),
            "globalRole": row[4],
            "createdAt": row[5],
        }

    def accounts(self) -> list[dict[str, Any]]:
        rows = self.database.execute(
            """
            SELECT id, email, name, enabled, global_role, created_at
            FROM platform_users ORDER BY created_at ASC, email ASC
            """
        ).fetchall()
        return [
            {
                "id": row[0],
                "email": row[1],
                "name": row[2],
                "enabled": bool(row[3]),
                "globalRole": row[4],
                "createdAt": row[5],
            }
            for row in rows
        ]

    def workspace_members(self, workspace_id: str) -> list[dict[str, Any]]:
        self._workspace_exists(workspace_id)
        rows = self.database.execute(
            """
            SELECT u.id, u.email, u.name, u.enabled, u.global_role, m.role,
                   u.created_at
            FROM workspace_members m
            JOIN platform_users u ON u.id = m.user_id
            WHERE m.workspace_id = ?
            ORDER BY m.role ASC, u.email ASC
            """,
            (workspace_id,),
        ).fetchall()
        return [
            {
                "id": row[0],
                "email": row[1],
                "name": row[2],
                "enabled": bool(row[3]),
                "globalRole": row[4],
                "role": normalize_workspace_role(row[5]),
                "createdAt": row[6],
            }
            for row in rows
            if is_workspace_role(normalize_workspace_role(row[5]))
        ]

    def _workspace_ids_for_user(self, user_id: str) -> list[str]:
        return [
            str(row[0])
            for row in self.database.execute(
                "SELECT workspace_id FROM workspace_members WHERE user_id = ?",
                (user_id,),
            ).fetchall()
        ]

    def _audit_account_change(
        self,
        actor_id: str,
        action: str,
        target_user_id: str,
        detail: dict[str, Any],
        workspace_ids: list[str] | None = None,
    ) -> None:
        targets = workspace_ids or self._workspace_ids_for_user(target_user_id)
        if not targets:
            targets = self._workspace_ids_for_user(actor_id)
        if not targets:
            first_workspace = self.database.execute(
                "SELECT id FROM workspaces ORDER BY created_at ASC LIMIT 1"
            ).fetchone()
            targets = [str(first_workspace[0])] if first_workspace else []
        for workspace_id in dict.fromkeys(targets):
            self.audit(
                workspace_id,
                {"type": "user", "id": actor_id},
                action,
                {"type": "user", "id": target_user_id},
                detail,
            )

    def _assert_not_last_workspace_admin(
        self, workspace_id: str, user_id: str
    ) -> None:
        from .http import PlatformError

        membership = self.database.execute(
            """
            SELECT m.role, u.enabled
            FROM workspace_members m
            JOIN platform_users u ON u.id = m.user_id
            WHERE m.workspace_id = ? AND m.user_id = ?
            """,
            (workspace_id, user_id),
        ).fetchone()
        if not membership or normalize_workspace_role(membership[0]) != WORKSPACE_ROLE_ADMIN:
            return
        if not bool(membership[1]):
            return
        count = self.database.execute(
            """
            SELECT COUNT(*)
            FROM workspace_members m
            JOIN platform_users u ON u.id = m.user_id
            WHERE m.workspace_id = ? AND m.role = ? AND u.enabled = 1
            """,
            (workspace_id, WORKSPACE_ROLE_ADMIN),
        ).fetchone()[0]
        if count <= 1:
            raise PlatformError(409, "LAST_WORKSPACE_ADMIN_REQUIRED")

    def _assert_not_last_super_admin(self, user_id: str) -> None:
        from .http import PlatformError

        row = self.database.execute(
            "SELECT global_role, enabled FROM platform_users WHERE id = ?",
            (user_id,),
        ).fetchone()
        if not row or not is_super_admin(row[0]) or not bool(row[1]):
            return
        count = self.database.execute(
            """
            SELECT COUNT(*) FROM platform_users
            WHERE global_role = ? AND enabled = 1
            """,
            (GLOBAL_ROLE_SUPER_ADMIN,),
        ).fetchone()[0]
        if count <= 1:
            raise PlatformError(409, "LAST_SUPER_ADMIN_REQUIRED")

    def _validate_invitation_email(self, email: str) -> str:
        from .http import PlatformError

        normalized = email.strip().lower()
        if not normalized or "@" not in normalized or len(normalized) > 320:
            raise PlatformError(400, "INVITE_EMAIL_INVALID")
        return normalized

    def invitation_response(self, row: tuple[Any, ...]) -> dict[str, Any]:
        expires_at = str(row[5])
        state = "active"
        if row[8]:
            state = "consumed"
        elif row[7]:
            state = "revoked"
        elif expires_at <= now():
            state = "expired"
        return {
            "id": row[0],
            "workspaceId": row[1],
            "email": row[2],
            "role": normalize_workspace_role(row[3]),
            "createdBy": row[4],
            "expiresAt": expires_at,
            "createdAt": row[6],
            "revokedAt": row[7],
            "consumedAt": row[8],
            "status": state,
        }

    def workspace_invitations(self, workspace_id: str) -> list[dict[str, Any]]:
        self._workspace_exists(workspace_id)
        rows = self.database.execute(
            """
            SELECT id, workspace_id, email, role, created_by, expires_at,
                   created_at, revoked_at, consumed_at
            FROM workspace_invitations
            WHERE workspace_id = ?
            ORDER BY created_at DESC
            """,
            (workspace_id,),
        ).fetchall()
        return [self.invitation_response(row) for row in rows]

    def create_workspace_invitation(
        self,
        workspace_id: str,
        actor_id: str,
        email: str,
        role: str,
    ) -> dict[str, Any]:
        from .http import PlatformError
        import secrets

        self._workspace_exists(workspace_id)
        email = self._validate_invitation_email(email)
        if not is_workspace_role(role):
            raise PlatformError(400, "INVITE_ROLE_INVALID")
        token = secrets.token_urlsafe(32)
        invitation = {
            "id": str(uuid.uuid4()),
            "workspaceId": workspace_id,
            "email": email,
            "role": role,
            "tokenHash": digest(token),
            "expiresAt": _iso_add_seconds(24 * 60 * 60),
            "createdAt": now(),
        }
        self.database.execute("BEGIN IMMEDIATE")
        try:
            replaced_at = now()
            replaced = self.database.execute(
                """
                SELECT id FROM workspace_invitations
                WHERE workspace_id = ? AND email = ? AND consumed_at IS NULL
                  AND revoked_at IS NULL AND expires_at > ?
                """,
                (workspace_id, email, replaced_at),
            ).fetchall()
            self.database.execute(
                """
                UPDATE workspace_invitations
                SET revoked_at = ?
                WHERE workspace_id = ? AND email = ? AND consumed_at IS NULL
                  AND revoked_at IS NULL AND expires_at > ?
                """,
                (replaced_at, workspace_id, email, replaced_at),
            )
            for previous in replaced:
                self.audit(
                    workspace_id,
                    {"type": "user", "id": actor_id},
                    "workspace.invitation_revoked",
                    {"type": "invitation", "id": str(previous[0])},
                    {"reason": "replaced"},
                )
            self.database.execute(
                """
                INSERT INTO workspace_invitations (
                  id, workspace_id, email, role, token_hash, expires_at,
                  created_by, created_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    invitation["id"],
                    workspace_id,
                    email,
                    role,
                    invitation["tokenHash"],
                    invitation["expiresAt"],
                    actor_id,
                    invitation["createdAt"],
                ),
            )
            self.audit(
                workspace_id,
                {"type": "user", "id": actor_id},
                "workspace.invitation_created",
                {"type": "invitation", "id": invitation["id"]},
                {"role": role},
            )
            self.database.execute("COMMIT")
        except Exception:
            self.database.execute("ROLLBACK")
            raise
        return {
            "id": invitation["id"],
            "workspaceId": workspace_id,
            "email": email,
            "role": role,
            "expiresAt": invitation["expiresAt"],
            "token": token,
        }

    def revoke_workspace_invitation(
        self, workspace_id: str, invitation_id: str, actor_id: str
    ) -> None:
        from .http import PlatformError

        self.database.execute("BEGIN IMMEDIATE")
        try:
            row = self.database.execute(
                """
                SELECT id, consumed_at, revoked_at, expires_at
                FROM workspace_invitations WHERE id = ? AND workspace_id = ?
                """,
                (invitation_id, workspace_id),
            ).fetchone()
            if not row:
                raise PlatformError(404, "INVITE_NOT_FOUND")
            if row[1] or row[2] or str(row[3]) <= now():
                raise PlatformError(409, "INVITE_NOT_REVOCABLE")
            self.database.execute(
                "UPDATE workspace_invitations SET revoked_at = ? WHERE id = ?",
                (now(), invitation_id),
            )
            self.audit(
                workspace_id,
                {"type": "user", "id": actor_id},
                "workspace.invitation_revoked",
                {"type": "invitation", "id": invitation_id},
                {},
            )
            self.database.execute("COMMIT")
        except Exception:
            self.database.execute("ROLLBACK")
            raise

    def accept_workspace_invitation(
        self,
        token: str,
        email: str,
        password: str | None,
        name: str | None,
        current_user: AuthUser | None = None,
    ) -> tuple[AuthUser, bool]:
        from .auth import password_hash
        from .http import PlatformError

        if not token or len(token) > 1024:
            raise PlatformError(400, "INVITE_TOKEN_INVALID")
        self.database.execute("BEGIN IMMEDIATE")
        try:
            invite = self.database.execute(
                """
                SELECT id, workspace_id, email, role, expires_at, revoked_at, consumed_at
                FROM workspace_invitations WHERE token_hash = ?
                """,
                (digest(token),),
            ).fetchone()
            if not invite:
                raise PlatformError(404, "INVITE_INVALID")
            # The consumed state deliberately wins before every other check so
            # replays have one terminal, non-enumerating contract.
            if invite[6]:
                raise PlatformError(410, "INVITE_ALREADY_USED")
            if invite[5]:
                raise PlatformError(410, "INVITE_REVOKED")
            if str(invite[4]) <= now():
                raise PlatformError(410, "INVITE_EXPIRED")
            # Do not validate caller-controlled fields before the consumed
            # check above. A replay must preserve its single terminal result
            # even when the caller changes or omits every other field.
            email = self._validate_invitation_email(email)
            if email != invite[2]:
                raise PlatformError(403, "INVITE_EMAIL_MISMATCH")
            role = normalize_workspace_role(invite[3])
            if not is_workspace_role(role):
                raise PlatformError(409, "INVITE_ROLE_INVALID")

            existing = self.database.execute(
                """
                SELECT id, email, name, enabled, global_role
                FROM platform_users WHERE email = ?
                """,
                (email,),
            ).fetchone()
            created = False
            if existing:
                # A disabled target cannot authenticate, so checking for its
                # matching session first would make the stable disabled-account
                # terminal response unreachable through the HTTP endpoint.
                if not bool(existing[3]):
                    raise PlatformError(403, "INVITE_ACCOUNT_DISABLED")
                if not current_user or current_user.id != existing[0]:
                    raise PlatformError(409, "INVITE_LOGIN_REQUIRED")
                user = AuthUser(existing[0], existing[1], existing[2], existing[4])
            else:
                if current_user:
                    raise PlatformError(403, "INVITE_EMAIL_MISMATCH")
                if not password or len(password) < 8 or len(password) > 1024:
                    raise PlatformError(400, "INVITE_PASSWORD_INVALID")
                created = True
                user = AuthUser(
                    str(uuid.uuid4()),
                    email,
                    (name or "").strip()[:100] or email.split("@", 1)[0],
                )
                created_at = now()
                self.database.execute(
                    """
                    INSERT INTO platform_users (id, email, name, created_at)
                    VALUES (?, ?, ?, ?)
                    """,
                    (user.id, user.email, user.name, created_at),
                )
                self.database.execute(
                    """
                    INSERT INTO platform_user_credentials
                      (user_id, password_hash, created_at, updated_at)
                    VALUES (?, ?, ?, ?)
                    """,
                    (user.id, password_hash(password), created_at, created_at),
                )

            membership = self.database.execute(
                """
                SELECT 1 FROM workspace_members
                WHERE workspace_id = ? AND user_id = ?
                """,
                (invite[1], user.id),
            ).fetchone()
            if membership:
                raise PlatformError(409, "INVITE_MEMBERSHIP_EXISTS")
            self.database.execute(
                """
                INSERT INTO workspace_members (workspace_id, user_id, role)
                VALUES (?, ?, ?)
                """,
                (invite[1], user.id, role),
            )
            cursor = self.database.execute(
                """
                UPDATE workspace_invitations SET consumed_at = ?
                WHERE id = ? AND consumed_at IS NULL
                """,
                (now(), invite[0]),
            )
            if cursor.rowcount != 1:
                raise PlatformError(410, "INVITE_ALREADY_USED")
            self.audit(
                invite[1],
                {"type": "user", "id": user.id},
                "workspace.invitation_accepted",
                {"type": "invitation", "id": invite[0]},
                {"role": role, "newAccount": created},
            )
            self.database.execute("COMMIT")
            return user, created
        except Exception:
            self.database.execute("ROLLBACK")
            raise

    def update_workspace_member_role(
        self, workspace_id: str, user_id: str, role: str, actor_id: str
    ) -> dict[str, Any]:
        from .http import PlatformError

        if not is_workspace_role(role):
            raise PlatformError(400, "WORKSPACE_ROLE_INVALID")
        self.database.execute("BEGIN IMMEDIATE")
        try:
            current = self.database.execute(
                """
                SELECT role FROM workspace_members
                WHERE workspace_id = ? AND user_id = ?
                """,
                (workspace_id, user_id),
            ).fetchone()
            if not current:
                raise PlatformError(404, "WORKSPACE_MEMBER_NOT_FOUND")
            current_role = normalize_workspace_role(current[0])
            if current_role == role:
                self.database.execute("COMMIT")
                return next(
                    item
                    for item in self.workspace_members(workspace_id)
                    if item["id"] == user_id
                )
            if current_role == WORKSPACE_ROLE_ADMIN and role != WORKSPACE_ROLE_ADMIN:
                self._assert_not_last_workspace_admin(workspace_id, user_id)
            self.database.execute(
                """
                UPDATE workspace_members SET role = ?
                WHERE workspace_id = ? AND user_id = ?
                """,
                (role, workspace_id, user_id),
            )
            revoked = self.revoke_user_sessions(user_id)
            self.audit(
                workspace_id,
                {"type": "user", "id": actor_id},
                "workspace.member_role_changed",
                {"type": "user", "id": user_id},
                {"role": role, "revokedSessions": revoked},
            )
            result = self.account_for(user_id)
            self.database.execute("COMMIT")
            return {**result, "role": role}
        except Exception:
            self.database.execute("ROLLBACK")
            raise

    def remove_workspace_member(
        self, workspace_id: str, user_id: str, actor_id: str
    ) -> None:
        from .http import PlatformError

        self.database.execute("BEGIN IMMEDIATE")
        try:
            membership = self.database.execute(
                """
                SELECT 1 FROM workspace_members
                WHERE workspace_id = ? AND user_id = ?
                """,
                (workspace_id, user_id),
            ).fetchone()
            if not membership:
                raise PlatformError(404, "WORKSPACE_MEMBER_NOT_FOUND")
            self._assert_not_last_workspace_admin(workspace_id, user_id)
            self.database.execute(
                "DELETE FROM workspace_members WHERE workspace_id = ? AND user_id = ?",
                (workspace_id, user_id),
            )
            revoked = self.revoke_user_sessions(user_id)
            self.audit(
                workspace_id,
                {"type": "user", "id": actor_id},
                "workspace.member_removed",
                {"type": "user", "id": user_id},
                {"revokedSessions": revoked},
            )
            self.database.execute("COMMIT")
        except Exception:
            self.database.execute("ROLLBACK")
            raise

    def set_account_enabled(
        self, user_id: str, enabled: bool, actor_id: str
    ) -> dict[str, Any]:
        self.database.execute("BEGIN IMMEDIATE")
        try:
            account = self.account_for(user_id)
            if account["enabled"] == enabled:
                self.database.execute("COMMIT")
                return account
            if not enabled:
                self._assert_not_last_super_admin(user_id)
                for workspace_id in self._workspace_ids_for_user(user_id):
                    self._assert_not_last_workspace_admin(workspace_id, user_id)
            self.database.execute(
                "UPDATE platform_users SET enabled = ? WHERE id = ?",
                (1 if enabled else 0, user_id),
            )
            revoked = self.revoke_user_sessions(user_id)
            self._audit_account_change(
                actor_id,
                "account.enabled_changed",
                user_id,
                {"enabled": enabled, "revokedSessions": revoked},
            )
            result = self.account_for(user_id)
            self.database.execute("COMMIT")
            return result
        except Exception:
            self.database.execute("ROLLBACK")
            raise

    def set_account_global_role(
        self, user_id: str, global_role: str | None, actor_id: str
    ) -> dict[str, Any]:
        from .http import PlatformError

        if global_role not in (None, GLOBAL_ROLE_SUPER_ADMIN):
            raise PlatformError(400, "GLOBAL_ROLE_INVALID")
        self.database.execute("BEGIN IMMEDIATE")
        try:
            account = self.account_for(user_id)
            if account["globalRole"] == global_role:
                self.database.execute("COMMIT")
                return account
            if global_role == GLOBAL_ROLE_SUPER_ADMIN and not account["enabled"]:
                raise PlatformError(409, "ACCOUNT_DISABLED")
            if account["globalRole"] == GLOBAL_ROLE_SUPER_ADMIN:
                self._assert_not_last_super_admin(user_id)
            self.database.execute(
                "UPDATE platform_users SET global_role = ? WHERE id = ?",
                (global_role, user_id),
            )
            revoked = self.revoke_user_sessions(user_id)
            self._audit_account_change(
                actor_id,
                "account.global_role_changed",
                user_id,
                {"globalRole": global_role, "revokedSessions": revoked},
            )
            result = self.account_for(user_id)
            self.database.execute("COMMIT")
            return result
        except Exception:
            self.database.execute("ROLLBACK")
            raise

    def issue_password_reset(self, user_id: str, actor_id: str) -> dict[str, Any]:
        import secrets

        self.database.execute("BEGIN IMMEDIATE")
        try:
            self.account_for(user_id)
            token = secrets.token_urlsafe(32)
            reset = {
                "id": str(uuid.uuid4()),
                "tokenHash": digest(token),
                "expiresAt": _iso_add_seconds(24 * 60 * 60),
                "createdAt": now(),
            }
            self.database.execute(
                """
                UPDATE password_reset_tokens SET revoked_at = ?
                WHERE user_id = ? AND consumed_at IS NULL AND revoked_at IS NULL
                  AND expires_at > ?
                """,
                (now(), user_id, now()),
            )
            self.database.execute(
                """
                INSERT INTO password_reset_tokens (
                  id, user_id, token_hash, expires_at, created_by, created_at
                ) VALUES (?, ?, ?, ?, ?, ?)
                """,
                (
                    reset["id"],
                    user_id,
                    reset["tokenHash"],
                    reset["expiresAt"],
                    actor_id,
                    reset["createdAt"],
                ),
            )
            self._audit_account_change(
                actor_id,
                "account.password_reset_issued",
                user_id,
                {},
            )
            self.database.execute("COMMIT")
            return {"id": reset["id"], "token": token, "expiresAt": reset["expiresAt"]}
        except Exception:
            self.database.execute("ROLLBACK")
            raise

    def accept_password_reset(self, token: str, password: str) -> None:
        from .auth import password_hash
        from .http import PlatformError

        if not token or len(token) > 1024:
            raise PlatformError(400, "PASSWORD_RESET_TOKEN_INVALID")
        self.database.execute("BEGIN IMMEDIATE")
        try:
            reset = self.database.execute(
                """
                SELECT id, user_id, expires_at, revoked_at, consumed_at
                FROM password_reset_tokens WHERE token_hash = ?
                """,
                (digest(token),),
            ).fetchone()
            if not reset:
                raise PlatformError(404, "PASSWORD_RESET_INVALID")
            if reset[4]:
                raise PlatformError(410, "PASSWORD_RESET_ALREADY_USED")
            if reset[3]:
                raise PlatformError(410, "PASSWORD_RESET_REVOKED")
            if str(reset[2]) <= now():
                raise PlatformError(410, "PASSWORD_RESET_EXPIRED")
            if not password or len(password) < 8 or len(password) > 1024:
                raise PlatformError(400, "PASSWORD_RESET_PASSWORD_INVALID")
            account = self.account_for(reset[1])
            if not account["enabled"]:
                raise PlatformError(403, "ACCOUNT_DISABLED")
            changed_at = now()
            self.database.execute(
                """
                UPDATE platform_user_credentials
                SET password_hash = ?, updated_at = ? WHERE user_id = ?
                """,
                (password_hash(password), changed_at, reset[1]),
            )
            self.database.execute(
                "UPDATE password_reset_tokens SET consumed_at = ? WHERE id = ?",
                (changed_at, reset[0]),
            )
            revoked = self.revoke_user_sessions(reset[1])
            self._audit_account_change(
                reset[1],
                "account.password_reset_accepted",
                reset[1],
                {"revokedSessions": revoked},
            )
            self.database.execute("COMMIT")
        except Exception:
            self.database.execute("ROLLBACK")
            raise

    def bootstrap_super_admin(
        self, email: str, name: str | None, password: str | None
    ) -> AuthUser:
        from .auth import password_hash
        from .http import PlatformError

        email = self._validate_invitation_email(email)
        self.database.execute("BEGIN IMMEDIATE")
        try:
            existing_admin = self.database.execute(
                """
                SELECT id FROM platform_users
                WHERE global_role = ? AND enabled = 1
                """,
                (GLOBAL_ROLE_SUPER_ADMIN,),
            ).fetchone()
            if existing_admin:
                raise PlatformError(409, "SUPER_ADMIN_ALREADY_CONFIGURED")
            existing = self.database.execute(
                """
                SELECT id, email, name, global_role FROM platform_users
                WHERE email = ?
                """,
                (email,),
            ).fetchone()
            created = False
            if existing:
                credential = self.database.execute(
                    "SELECT 1 FROM platform_user_credentials WHERE user_id = ?",
                    (existing[0],),
                ).fetchone()
                if not credential:
                    if not password or len(password) < 8 or len(password) > 1024:
                        raise PlatformError(400, "BOOTSTRAP_PASSWORD_INVALID")
                    created_at = now()
                    self.database.execute(
                        """
                        INSERT INTO platform_user_credentials
                          (user_id, password_hash, created_at, updated_at)
                        VALUES (?, ?, ?, ?)
                        """,
                        (existing[0], password_hash(password), created_at, created_at),
                    )
                self.database.execute(
                    """
                    UPDATE platform_users
                    SET enabled = 1, global_role = ? WHERE id = ?
                    """,
                    (GLOBAL_ROLE_SUPER_ADMIN, existing[0]),
                )
                user = AuthUser(existing[0], existing[1], existing[2], GLOBAL_ROLE_SUPER_ADMIN)
            else:
                if not password or len(password) < 8 or len(password) > 1024:
                    raise PlatformError(400, "BOOTSTRAP_PASSWORD_INVALID")
                user = AuthUser(
                    str(uuid.uuid4()),
                    email,
                    (name or "").strip()[:100] or email.split("@", 1)[0],
                    GLOBAL_ROLE_SUPER_ADMIN,
                )
                created = True
                created_at = now()
                self.database.execute(
                    """
                    INSERT INTO platform_users (id, email, name, global_role, created_at)
                    VALUES (?, ?, ?, ?, ?)
                    """,
                    (user.id, user.email, user.name, user.global_role, created_at),
                )
                self.database.execute(
                    """
                    INSERT INTO platform_user_credentials
                      (user_id, password_hash, created_at, updated_at)
                    VALUES (?, ?, ?, ?)
                    """,
                    (user.id, password_hash(password), created_at, created_at),
                )
            revoked_sessions = self.revoke_user_sessions(user.id)
            self.deployment_audit(
                {"type": "system", "id": "bootstrap"},
                "account.super_admin_bootstrapped",
                {"type": "user", "id": user.id},
                {
                    "created": created,
                    "promoted": not created,
                    "revokedSessions": revoked_sessions,
                },
            )
            self.database.execute("COMMIT")
            return user
        except Exception:
            self.database.execute("ROLLBACK")
            raise

    def document_for(self, project_id: str) -> dict[str, Any]:
        row = self.database.execute(
            """
            SELECT data, version, updated_at FROM project_documents
            WHERE project_id = ?
            """,
            (project_id,),
        ).fetchone()
        if row:
            return {
                "data": parse_json(row[0], {}),
                "version": row[1],
                "updatedAt": row[2],
            }
        return {"data": {}, "version": 0, "updatedAt": None}

    def dataset_version_for(self, project_id: str, version_id: str) -> dict[str, Any]:
        from .http import PlatformError

        row = self.database.execute(
            """
            SELECT v.id, v.dataset_id, d.project_id, v.version_number,
                   v.columns_json, v.row_count, v.checksum, v.source_name,
                   v.created_at
            FROM dataset_versions v
            JOIN datasets d ON d.id = v.dataset_id
            WHERE v.id = ? AND d.project_id = ? AND d.archived_at IS NULL
            """,
            (version_id, project_id),
        ).fetchone()
        if not row:
            raise PlatformError(404, "DATASET_VERSION_NOT_FOUND")
        return {
            "id": row[0],
            "datasetId": row[1],
            "projectId": row[2],
            "versionNumber": row[3],
            "columns": parse_json(row[4], []),
            "rowCount": row[5],
            "checksum": row[6],
            "sourceName": row[7],
            "createdAt": row[8],
        }

    def dataset_version_response(self, row: tuple[Any, ...]) -> dict[str, Any]:
        return {
            "id": row[0],
            "datasetId": row[1],
            "projectId": row[2],
            "versionNumber": row[3],
            "columns": parse_json(row[4], []),
            "rowCount": row[5],
            "checksum": row[6],
            "sourceName": row[7],
            "createdAt": row[8],
        }

    def dataset_rows_for(self, version_id: str) -> list[dict[str, Any]]:
        rows = self.database.execute(
            """
            SELECT row_number, data_json FROM dataset_rows
            WHERE dataset_version_id = ? ORDER BY row_number ASC
            """,
            (version_id,),
        ).fetchall()
        return [
            {
                "rowNumber": row[0],
                "data": parse_json(row[1], {}),
            }
            for row in rows
        ]

    def published_revision_for(
        self,
        project_id: str,
        revision_id: str | None = None,
        *,
        flow_id: str | None = None,
        environment_id: str | None = None,
        allow_superseded: bool = False,
    ) -> dict[str, Any]:
        from .http import PlatformError

        flow_context_id = flow_id or None
        if revision_id:
            statuses = (
                ("published", "superseded") if allow_superseded else ("published",)
            )
            status_filter = ",".join("?" for _ in statuses)
            row = self.database.execute(
                f"""
                SELECT id, flow_snapshot, environment_snapshot,
                       element_snapshot, dataset_snapshot, checksum, flow_id
                FROM flow_revisions
                WHERE id = ? AND project_id = ?
                  AND status IN ({status_filter})
                """,
                (revision_id, project_id, *statuses),
            ).fetchone()
            if not row:
                raise PlatformError(409, "PUBLISHED_REVISION_REQUIRED")
            revision_flow_id = (
                row[6] if isinstance(row[6], str) and row[6] else None
            )
            if revision_flow_id is None:
                snapshot_flow = parse_json(row[1], {})
                snapshot_flow_id = (
                    snapshot_flow.get("id")
                    if isinstance(snapshot_flow, dict)
                    else None
                )
                revision_flow_id = (
                    snapshot_flow_id
                    if isinstance(snapshot_flow_id, str)
                    else None
                )
            if (
                flow_context_id
                and revision_flow_id
                and flow_context_id != revision_flow_id
            ):
                raise PlatformError(409, "REVISION_FLOW_MISMATCH")
        else:
            if not flow_context_id:
                raise PlatformError(400, "FLOW_ID_REQUIRED")
            environment_filter = ""
            query_params: list[Any] = [project_id, flow_context_id]
            if environment_id:
                environment_filter = " AND environment_id = ?"
                query_params.append(environment_id)
            row = self.database.execute(
                f"""
                SELECT id, flow_snapshot, environment_snapshot,
                       element_snapshot, dataset_snapshot, checksum, flow_id
                FROM flow_revisions
                WHERE project_id = ? AND flow_id = ? AND status = 'published'
                  {environment_filter}
                ORDER BY published_at DESC, revision_number DESC, created_at DESC
                LIMIT 1
                """,
                tuple(query_params),
            ).fetchone()
            if not row:
                raise PlatformError(409, "PUBLISHED_REVISION_REQUIRED")
        return {
            "id": row[0],
            "flow_snapshot": row[1],
            "environment_snapshot": row[2],
            "element_snapshot": row[3],
            "dataset_snapshot": row[4],
            "checksum": row[5],
            "flow_id": row[6],
        }

    def require_revision_environment(
        self, revision: dict[str, Any], environment_id: str
    ) -> None:
        from .http import PlatformError

        environment = parse_json(revision["environment_snapshot"], {})
        snapshot_id = environment.get("id") if isinstance(environment, dict) else ""
        if snapshot_id != environment_id:
            raise PlatformError(409, "REVISION_ENVIRONMENT_MISMATCH")

    def managed_agent(self, project_id: str) -> dict[str, Any]:
        project = self.project_for(project_id)
        agent_id = f"managed-{project['workspace_id']}"
        self.database.execute(
            """
            INSERT OR IGNORE INTO agents (
              id, workspace_id, name, credential_hash, status,
              browser_version, os, max_concurrency, created_at
            ) VALUES (?, ?, 'ManagedRunner', ?, 'disabled', 'bundled',
                      'Windows', 1, ?)
            """,
            (agent_id, project["workspace_id"], digest(f"managed:{project['workspace_id']}"), now()),
        )
        return {
            "id": agent_id,
            "workspaceId": project["workspace_id"],
            "name": "ManagedRunner",
            "status": "disabled",
            "browserVersion": "bundled",
            "os": "Windows",
            "maxConcurrency": 1,
            "currentTask": None,
            "lastSeenAt": None,
            "createdAt": now(),
        }

    def missing_secret_names(
        self, project_id: str, requested: list[str]
    ) -> list[str]:
        if not requested:
            return []
        placeholders = ",".join("?" for _ in requested)
        rows = self.database.execute(
            f"""
            SELECT name FROM project_secrets
            WHERE project_id = ? AND name IN ({placeholders})
            """,
            (project_id, *requested),
        ).fetchall()
        found = {row[0] for row in rows}
        return [name for name in requested if name not in found]

    def secret_values(self, project_id: str, requested: list[str]) -> dict[str, str]:
        from .http import PlatformError

        if not requested:
            return {}
        placeholders = ",".join("?" for _ in requested)
        rows = self.database.execute(
            f"""
            SELECT name, iv, tag, ciphertext FROM project_secrets
            WHERE project_id = ? AND name IN ({placeholders})
            """,
            (project_id, *requested),
        ).fetchall()
        if len(rows) != len(requested):
            raise PlatformError(409, "RUN_SECRET_NOT_CONFIGURED")
        project = self.project_for(project_id)
        self.audit(
            project["workspace_id"],
            {"type": "system", "id": "managed-runner"},
            "secret.decrypted_for_run",
            {"type": "project", "id": project_id},
            {"names": requested},
            project_id,
        )
        return {
            row[0]: self.decrypt({"iv": row[1], "tag": row[2], "ciphertext": row[3]})
            for row in rows
        }

    def notification_target(self, value: str) -> dict[str, str]:
        from urllib.parse import urlsplit

        from .core import (
            ALLOW_INSECURE_NOTIFICATION_TARGETS,
            ALLOW_PRIVATE_NOTIFICATION_TARGETS,
            NOTIFICATION_HOST_ALLOWLIST,
            notification_host_allowed_list,
            public_ip_address,
        )

        target = urlsplit(value)
        if target.username or target.password:
            raise ValueError("NOTIFICATION_URL_CREDENTIALS_FORBIDDEN")
        if target.scheme != "https" and not (
            ALLOW_INSECURE_NOTIFICATION_TARGETS and target.scheme == "http"
        ):
            raise ValueError("NOTIFICATION_URL_PROTOCOL_FORBIDDEN")
        host = (target.hostname or "").lower()
        explicitly_allowed = notification_host_allowed_list(host)
        if NOTIFICATION_HOST_ALLOWLIST and not explicitly_allowed:
            raise ValueError("NOTIFICATION_URL_HOST_NOT_ALLOWED")
        if (
            not ALLOW_PRIVATE_NOTIFICATION_TARGETS or not explicitly_allowed
        ) and (
            host == "localhost"
            or host.endswith(".localhost")
            or host.endswith(".local")
        ):
            raise ValueError("NOTIFICATION_URL_PRIVATE_HOST")
        try:
            ipaddress.ip_address(host)
            addresses = [host]
        except ValueError:
            addresses = [
                entry[4][0]
                for entry in socket.getaddrinfo(host, target.port or None)
            ]
        if (
            not ALLOW_PRIVATE_NOTIFICATION_TARGETS or not explicitly_allowed
        ) and (
            not addresses
            or any(not public_ip_address(address) for address in addresses)
        ):
            raise ValueError("NOTIFICATION_URL_PRIVATE_HOST")
        if not addresses:
            raise ValueError("NOTIFICATION_URL_HOST_UNRESOLVED")
        return {"url": value, "address": addresses[0]}

    def deliver_pending_notifications(self) -> None:
        from .core import (
            NOTIFICATION_MAX_ATTEMPTS,
            NOTIFICATION_RETRY_BASE_MS,
        )

        stale_claim = _iso_from_ms(_now_ms() - 30_000)
        stale_rows = self.database.execute(
            """
            SELECT id, attempt_count FROM deliveries
            WHERE status = 'delivering' AND updated_at <= ?
            """,
            (stale_claim,),
        ).fetchall()
        for row in stale_rows:
            attempts = int(row[1]) + 1
            delay_ms = NOTIFICATION_RETRY_BASE_MS * 2 ** max(0, attempts - 1)
            self.database.execute(
                """
                UPDATE deliveries
                SET status = 'retrying', next_attempt_at = ?, updated_at = ?
                WHERE id = ? AND status = 'delivering'
                """,
                (_iso_from_ms(_now_ms() + delay_ms), now(), row[0]),
            )

        due_rows = self.database.execute(
            """
            SELECT d.id, d.run_id, d.channel_id, d.payload, d.attempt_count,
                   c.channel_type, c.name AS channel_name, c.workspace_id,
                   c.config_iv, c.config_tag, c.config_ciphertext
            FROM deliveries d
            JOIN notification_channels c ON c.id = d.channel_id
            WHERE d.status IN ('pending', 'retrying') AND c.enabled = 1
              AND COALESCE(d.next_attempt_at, d.created_at) <= ?
            ORDER BY d.created_at ASC LIMIT 20
            """,
            (now(),),
        ).fetchall()
        for delivery in due_rows:
            claimed = self.database.execute(
                """
                UPDATE deliveries
                SET status = 'delivering', attempt_count = attempt_count + 1,
                    updated_at = ?
                WHERE id = ? AND status IN ('pending', 'retrying')
                """,
                (now(), delivery[0]),
            )
            if claimed.rowcount != 1:
                continue
            status = "failed"
            response_code: int | None = None
            error: str | None = None
            try:
                config = parse_json(
                    self.decrypt(
                        {
                            "iv": delivery[8],
                            "tag": delivery[9],
                            "ciphertext": delivery[10],
                        }
                    ),
                    {},
                )
                endpoint = config.get("url") if isinstance(config, dict) else None
                if not isinstance(endpoint, str):
                    raise ValueError("NOTIFICATION_CONFIG_INVALID")
                target = self.notification_target(endpoint)
                headers = as_record(config.get("headers"))
                string_headers = {
                    str(key): str(value)
                    for key, value in headers.items()
                    if isinstance(value, str)
                }
                keyword = config.get("keyword")
                keyword = (
                    keyword.strip()[:200]
                    if isinstance(keyword, str) and keyword.strip()
                    else None
                )
                response = _post_notification(
                    target,
                    string_headers,
                    json(
                        _format_notification_body(
                            delivery[5], parse_json(delivery[3], {}), keyword
                        )
                    ),
                )
                response_code = response["status"]
                status = (
                    "delivered"
                    if 200 <= response_code < 300
                    else "failed"
                )
                error = None if status == "delivered" else f"HTTP_{response_code}"
                if status == "delivered" and response["body"]:
                    body_code = notification_rejection_code(response["body"])
                    if body_code is not None:
                        status = "failed"
                        error = f"NOTIFICATION_REJECTED_{body_code}"
            except Exception as exc:
                error = (
                    "NOTIFICATION_TIMEOUT"
                    if isinstance(exc, TimeoutError)
                    else str(exc)[:200]
                )
            attempts = int(delivery[4]) + 1
            retry = status == "failed" and attempts < NOTIFICATION_MAX_ATTEMPTS
            next_attempt_at = (
                _iso_from_ms(
                    _now_ms()
                    + NOTIFICATION_RETRY_BASE_MS * 2 ** max(0, attempts - 1)
                )
                if retry
                else None
            )
            self.database.execute(
                """
                UPDATE deliveries
                SET status = ?, attempt_count = ?, response_code = ?,
                    error = ?, next_attempt_at = ?,
                    delivered_at = CASE WHEN ? = 'delivered' THEN ?
                      ELSE delivered_at END,
                    updated_at = ?
                WHERE id = ? AND status = 'delivering'
                """,
                (
                    "retrying" if retry else status,
                    attempts,
                    response_code,
                    error,
                    next_attempt_at,
                    status,
                    now(),
                    now(),
                    delivery[0],
                ),
            )
            if retry:
                continue
            delivery_project = self.database.execute(
                "SELECT project_id FROM platform_runs WHERE id = ?",
                (delivery[1],),
            ).fetchone()
            delivery_action = (
                "notification.delivered"
                if status == "delivered"
                else "notification.rejected"
                if error and error.startswith("NOTIFICATION_REJECTED_")
                else "notification.failed"
            )
            self.audit(
                delivery[7],
                {"type": "system", "id": f"delivery:{delivery[0]}"},
                delivery_action,
                {"type": "notification_channel", "id": delivery[2]},
                {
                    "channelType": delivery[5],
                    "channelName": delivery[6],
                    "code": response_code,
                    "error": (error or "")[:200] or None,
                },
                delivery_project[0] if delivery_project else None,
            )

    def send_test_notification(
        self, channel_id: str, workspace_id: str | None = None
    ) -> dict[str, Any]:
        from .http import PlatformError

        row = self.database.execute(
            """
            SELECT id, channel_type, config_iv, config_tag, config_ciphertext
            FROM notification_channels
            WHERE id = ? AND archived_at IS NULL
              AND (? IS NULL OR workspace_id = ?)
            """,
            (channel_id, workspace_id, workspace_id),
        ).fetchone()
        if not row:
            raise PlatformError(404, "NOTIFICATION_CHANNEL_NOT_FOUND")
        config = parse_json(
            self.decrypt(
                {
                    "iv": row[2],
                    "tag": row[3],
                    "ciphertext": row[4],
                }
            ),
            {},
        )
        endpoint = config.get("url") if isinstance(config, dict) else None
        if not isinstance(endpoint, str):
            raise PlatformError(400, "NOTIFICATION_CONFIG_INVALID")
        target = self.notification_target(endpoint)
        headers = as_record(config.get("headers"))
        string_headers = {
            str(key): str(value)
            for key, value in headers.items()
            if isinstance(value, str)
        }
        response = _post_notification(
            target,
            string_headers,
            json(
                {
                    "type": "test",
                    "message": "AutoFlow test notification",
                    "timestamp": now(),
                }
            ),
        )
        status = response["status"]
        error = None if 200 <= status < 300 else f"HTTP_{status}"
        if error is None and response.get("body"):
            body_code = notification_rejection_code(response["body"])
            if body_code is not None:
                error = f"NOTIFICATION_REJECTED_{body_code}"
        return {"status": status, "error": error}

    def append_run_event(
        self, run_id: str, kind: str, data: dict[str, Any]
    ) -> None:
        self.database.execute(
            """
            INSERT INTO platform_run_events (run_id, kind, data, created_at)
            VALUES (?, ?, ?, ?)
            """,
            (run_id, kind, json(data), now()),
        )

    def resolve_run_spec(self, input: dict[str, Any]) -> dict[str, Any]:
        from .http import PlatformError

        project_id = input["projectId"]
        revision = self.published_revision_for(
            project_id,
            input.get("revisionId") or None,
            flow_id=input.get("flowId") or None,
            environment_id=input.get("environmentId") or None,
            allow_superseded=bool(input.get("allowSuperseded")),
        )
        environment = parse_json(revision["environment_snapshot"], {})
        environment_id = input.get("environmentId")
        if not environment_id:
            environment_id = (
                environment.get("id") if isinstance(environment, dict) else ""
            )
        if not environment_id:
            raise PlatformError(400, "ENVIRONMENT_REQUIRED")
        if environment.get("id") != environment_id:
            raise PlatformError(409, "REVISION_ENVIRONMENT_MISMATCH")
        self.require_chromium_environment(
            environment if isinstance(environment, dict) else {}
        )
        self.ensure_chromium_available()
        agent = self.managed_agent(project_id)
        dataset_version_id = input.get("datasetVersionId")
        if not dataset_version_id:
            snapshot_dataset = parse_json(revision["dataset_snapshot"], None)
            if isinstance(snapshot_dataset, dict):
                version_id = snapshot_dataset.get("versionId")
                if isinstance(version_id, str):
                    dataset_version_id = version_id
        dataset_version = (
            self.dataset_version_for(project_id, dataset_version_id)
            if dataset_version_id
            else None
        )
        rows = (
            self.dataset_rows_for(dataset_version["id"])
            if dataset_version
            else [{"rowNumber": None, "data": None}]
        )
        max_runs = input.get("maxRuns")
        if max_runs is not None and len(rows) > int(max_runs):
            raise PlatformError(413, "RUN_COUNT_LIMIT_EXCEEDED")
        flow = parse_json(revision["flow_snapshot"], {})
        flow_steps = flow.get("steps", []) if isinstance(flow, dict) else []
        if not isinstance(flow_steps, list):
            flow_steps = []
        up_to_step_id = input.get("upToStepId")
        if up_to_step_id and not any(
            as_record(step).get("id") == up_to_step_id for step in flow_steps
        ):
            raise PlatformError(400, "RUN_STEP_NOT_FOUND")
        secret_names = flow.get("secretNames", [])
        if not isinstance(secret_names, list):
            secret_names = []
        secret_names = [value for value in secret_names if isinstance(value, str)]
        step_limit = (
            next(
                (
                    index + 1
                    for index, step in enumerate(flow_steps)
                    if as_record(step).get("id") == up_to_step_id
                ),
                0,
            )
            if up_to_step_id
            else len(flow_steps)
        )
        required_secret_names: set[str] = set()
        for step in flow_steps[:step_limit]:
            value = as_record(step).get("value")
            if not isinstance(value, str):
                value = ""
            for name in secret_names:
                if (
                    f"{{{{{name}}}}}" in value
                    or f"{{{{ {name} }}}}" in value
                    or f"{{{{secret.{name}}}}}" in value
                    or f"{{{{ secret.{name} }}}}" in value
                ):
                    required_secret_names.add(name)
        return {
            "projectId": project_id,
            "revision": revision,
            "environment": environment,
            "environmentId": environment_id,
            "datasetVersionId": dataset_version_id,
            "datasetVersion": dataset_version,
            "rows": rows,
            "flow": flow,
            "flowSteps": flow_steps,
            "secretNames": secret_names,
            "requiredSecretNames": required_secret_names,
            "upToStepId": up_to_step_id or None,
            "agent": agent,
        }

    @staticmethod
    def _required_retry_variable_names(
        flow: dict[str, Any], secret_names: list[str]
    ) -> set[str]:
        secret_set = {str(name) for name in secret_names}
        required: set[str] = set()
        steps = flow.get("steps") if isinstance(flow, dict) else []
        if not isinstance(steps, list):
            steps = []
        for step in steps:
            value = as_record(step).get("value")
            if not isinstance(value, str):
                continue
            for match in re.finditer(r"{{\s*([^}]+)\s*}}", value):
                expression = match.group(1).strip()
                if not expression:
                    continue
                if (
                    expression.startswith("secret.")
                    or expression in secret_set
                    or expression.startswith("data.")
                    or expression.startswith("flow.")
                    or expression.startswith("run.")
                    or expression == "env.baseUrl"
                ):
                    continue
                required.add(expression)
        return required

    def _preflight_retry_variables(
        self, project_id: str, flow: dict[str, Any], secret_names: list[str]
    ) -> None:
        from .http import PlatformError

        required = self._required_retry_variable_names(flow, secret_names)
        if not required:
            return
        rows = self.database.execute(
            """
            SELECT data FROM project_resources
            WHERE project_id = ? AND resource_type = 'variables'
              AND archived_at IS NULL
            """,
            (project_id,),
        ).fetchall()
        available: set[str] = set()
        for variable_row in rows:
            variable = parse_json(variable_row[0], {})
            if (
                not isinstance(variable, dict)
                or variable.get("secret") is True
                or not isinstance(variable.get("name"), str)
                or not isinstance(variable.get("value"), str)
            ):
                continue
            scope = (
                "env"
                if variable.get("scope") == "环境"
                else "project"
                if variable.get("scope") == "项目"
                else ""
            )
            key = f"{scope}.{variable['name']}" if scope else variable["name"]
            available.add(key)
        missing = sorted(required - available)
        if missing:
            raise PlatformError(409, "RUN_VARIABLE_NOT_CONFIGURED")

    def _preflight_retry_spec(self, project_id: str, spec: dict[str, Any]) -> None:
        from .http import PlatformError

        snapshot_secret_names = [
            name
            for name in spec["secretNames"]
            if name in spec["requiredSecretNames"]
        ]
        missing = self.missing_secret_names(project_id, snapshot_secret_names)
        if missing:
            raise PlatformError(409, "RUN_SECRET_NOT_CONFIGURED")
        self._preflight_retry_variables(
            project_id, spec["flow"], spec["secretNames"]
        )

    @staticmethod
    def _row_from_retry_snapshot(snapshot: dict[str, Any]) -> dict[str, Any]:
        dataset_row = snapshot.get("datasetRow")
        if isinstance(dataset_row, dict):
            return {
                "rowNumber": dataset_row.get("number"),
                "data": dataset_row.get("data"),
            }
        return {"rowNumber": None, "data": None}

    def clone_retry_spec_from_run(self, run: dict[str, Any]) -> dict[str, Any]:
        """从源 run 的持久 snapshot 构造一对一 retry spec，不重新解析 revision/dataset。"""
        from .http import PlatformError

        snapshot = run.get("snapshot")
        if not isinstance(snapshot, dict):
            raise PlatformError(409, "RUN_SNAPSHOT_NOT_RETRYABLE")
        flow = snapshot.get("flow")
        environment = snapshot.get("environment")
        if not isinstance(flow, dict) or not isinstance(environment, dict):
            raise PlatformError(409, "RUN_SNAPSHOT_NOT_RETRYABLE")
        revision_id = snapshot.get("flowRevisionId")
        revision_checksum = snapshot.get("flowRevisionChecksum")
        if (
            not isinstance(revision_id, str)
            or not isinstance(revision_checksum, str)
            or run.get("revisionId") != revision_id
        ):
            raise PlatformError(409, "RUN_SNAPSHOT_NOT_RETRYABLE")
        if environment.get("id") != run.get("environmentId"):
            raise PlatformError(409, "RUN_SNAPSHOT_NOT_RETRYABLE")
        up_to_step_id = snapshot.get("upToStepId")
        steps = flow.get("steps") if isinstance(flow.get("steps"), list) else []
        if up_to_step_id and not any(
            as_record(step).get("id") == up_to_step_id for step in steps
        ):
            raise PlatformError(409, "RUN_SNAPSHOT_NOT_RETRYABLE")
        elements = snapshot.get("elements")
        if not isinstance(elements, list):
            raise PlatformError(409, "RUN_SNAPSHOT_NOT_RETRYABLE")
        dataset = snapshot.get("dataset")
        if dataset is not None and (
            not isinstance(dataset, dict)
            or not all(
                key in dataset
                for key in ("datasetId", "versionId", "versionNumber", "checksum", "columns")
            )
        ):
            raise PlatformError(409, "RUN_SNAPSHOT_NOT_RETRYABLE")
        dataset_row = snapshot.get("datasetRow")
        if dataset_row is not None and (
            not isinstance(dataset_row, dict)
            or "number" not in dataset_row
            or "data" not in dataset_row
        ):
            raise PlatformError(409, "RUN_SNAPSHOT_NOT_RETRYABLE")
        secret_names = snapshot.get("secretNames", [])
        if not isinstance(secret_names, list) or not all(
            isinstance(name, str) for name in secret_names
        ):
            raise PlatformError(409, "RUN_SNAPSHOT_NOT_RETRYABLE")
        return {
            "projectId": run["projectId"],
            "revision": {
                "id": revision_id,
                "checksum": revision_checksum,
                "element_snapshot": json(elements),
            },
            "environmentId": run["environmentId"],
            "environment": _json.loads(_json.dumps(environment)),
            "flow": _json.loads(_json.dumps(flow)),
            "flowSteps": _json.loads(_json.dumps(steps)),
            "datasetVersionId": (
                dataset.get("versionId") if isinstance(dataset, dict) else None
            ),
            "datasetVersion": (
                {
                    "id": dataset["versionId"],
                    "datasetId": dataset["datasetId"],
                    "versionNumber": dataset["versionNumber"],
                    "checksum": dataset["checksum"],
                    "columns": dataset["columns"],
                }
                if isinstance(dataset, dict)
                else None
            ),
            "secretNames": list(secret_names),
            "requiredSecretNames": set(secret_names),
            "upToStepId": up_to_step_id or None,
            "agent": self.managed_agent(run["projectId"]),
        }

    def retry_run_snapshot(
        self, project_id: str, run_id: str, actor_id: str
    ) -> dict[str, Any]:
        from .http import PlatformError

        run = self.run_by_id(run_id)
        if run["projectId"] != project_id:
            raise PlatformError(404, "RUN_NOT_FOUND")
        if run["status"] not in ("failed", "canceled"):
            raise PlatformError(409, "RUN_NOT_RETRYABLE")
        spec = self.clone_retry_spec_from_run(run)
        self._preflight_retry_spec(project_id, spec)
        row = self._row_from_retry_snapshot(run["snapshot"])
        new_run_id: str | None = None
        self.database.execute("BEGIN IMMEDIATE")
        try:
            new_run_id = self.insert_run_from_spec(
                spec,
                row=row,
                created_by=actor_id,
                source="manual",
                retry_of_run_id=run_id,
            )
            self.database.execute("COMMIT")
        except Exception:
            self.database.execute("ROLLBACK")
            raise
        if new_run_id is None:
            raise RuntimeError("retry_run_snapshot failed to create run")
        self.enqueue_managed_run(new_run_id)
        return {
            "runIds": [new_run_id],
            "runs": [self.run_response(self.run_by_id(new_run_id))],
        }

    def insert_run_from_spec(
        self,
        spec: dict[str, Any],
        *,
        row: dict[str, Any],
        created_by: str,
        source: str,
        dispatch_key: str | None = None,
        batch_id: str | None = None,
        batch_item_index: int | None = None,
        retry_of_run_id: str | None = None,
    ) -> str:
        from .http import PlatformError

        project_id = spec["projectId"]
        revision = spec["revision"]
        environment_id = spec["environmentId"]
        environment = spec["environment"]
        dataset_version = spec["datasetVersion"]
        dataset_version_id = spec["datasetVersionId"]
        flow = spec["flow"]
        agent = spec["agent"]
        up_to_step_id = spec["upToStepId"]
        if dispatch_key:
            existing = self.database.execute(
                "SELECT id FROM platform_runs WHERE dispatch_key = ?",
                (dispatch_key,),
            ).fetchone()
            if existing:
                return existing[0]
        snapshot_secret_names = [
            name
            for name in spec["secretNames"]
            if name in spec["requiredSecretNames"]
        ]
        missing = self.missing_secret_names(project_id, snapshot_secret_names)
        if missing:
            raise PlatformError(409, "RUN_SECRET_NOT_CONFIGURED")
        run_id = str(uuid.uuid4())
        created_at = now()
        # 运行 snapshot 不保存普通变量值；变量在 enqueue/恢复时按项目当前值读取。
        flow_snapshot = dict(flow)
        flow_snapshot.pop("variables", None)
        snapshot = {
            "flowRevisionId": revision["id"],
            "flowRevisionChecksum": revision["checksum"],
            "environmentId": environment_id,
            "flow": flow_snapshot,
            "environment": environment,
            "elements": parse_json(revision["element_snapshot"], []),
            "dataset": (
                {
                    "datasetId": dataset_version["datasetId"],
                    "versionId": dataset_version["id"],
                    "versionNumber": dataset_version["versionNumber"],
                    "checksum": dataset_version["checksum"],
                    "columns": dataset_version["columns"],
                }
                if dataset_version
                else None
            ),
            "datasetRow": (
                {"number": row["rowNumber"], "data": row["data"]}
                if row["data"] is not None
                else None
            ),
            "secretNames": snapshot_secret_names,
            "upToStepId": up_to_step_id,
            "executor": {
                "type": "managed",
                "id": agent["id"],
                "name": agent["name"],
                "browserVersion": agent["browserVersion"],
            },
            "trigger": source,
        }
        if batch_id:
            snapshot["batchId"] = batch_id
            snapshot["batchItemIndex"] = batch_item_index
        self.database.execute(
            """
            INSERT INTO platform_runs (
              id, project_id, revision_id, environment_id, agent_id,
              executor_type, dispatch_key, status, snapshot,
              batch_id, batch_item_index, retry_of_run_id,
              created_by, created_at, updated_at
            ) VALUES (?, ?, ?, ?, ?, 'managed', ?, 'queued', ?, ?, ?, ?, ?, ?, ?)
            """,
            (
                run_id,
                project_id,
                revision["id"],
                environment_id,
                agent["id"],
                dispatch_key,
                json(snapshot),
                batch_id,
                batch_item_index,
                retry_of_run_id,
                created_by,
                created_at,
                created_at,
            ),
        )
        self.append_run_event(
            run_id,
            "run.queued",
            {
                "revisionId": revision["id"],
                "environmentId": environment_id,
                "executorType": "managed",
                "source": source,
                "datasetVersionId": dataset_version_id,
                "datasetRow": row["rowNumber"],
            },
        )
        if retry_of_run_id:
            self.append_run_event(
                run_id,
                "run.retried",
                {"priorRunId": retry_of_run_id, "actorId": created_by},
            )
        return run_id

    def queue_published_runs(self, input: dict[str, Any]) -> dict[str, Any]:
        spec = self.resolve_run_spec(input)
        run_ids: list[str] = []
        self.database.execute("BEGIN IMMEDIATE")
        try:
            for row in spec["rows"]:
                dispatch_key = (
                    f"{input['dispatchKey']}:{row['rowNumber'] or 0}"
                    if input.get("dispatchKey")
                    else None
                )
                run_ids.append(
                    self.insert_run_from_spec(
                        spec,
                        row=row,
                        created_by=input["createdBy"],
                        source=input["source"],
                        dispatch_key=dispatch_key,
                    )
                )
            self.database.execute("COMMIT")
        except Exception:
            self.database.execute("ROLLBACK")
            raise
        for run_id in run_ids:
            self.enqueue_managed_run(run_id)
        return {
            "runIds": run_ids,
            "revision": spec["revision"],
            "environmentId": spec["environmentId"],
            "datasetVersionId": spec["datasetVersionId"],
        }

    BATCH_MIN_FLOWS = 2
    BATCH_MAX_FLOWS = 20
    BATCH_MAX_TOTAL_STEPS = 2000

    _RUN_BATCH_COUNTS_CTE = """
        counts AS (
          SELECT batch_id,
                 COUNT(*) AS total,
                 SUM(CASE WHEN status = 'queued' THEN 1 ELSE 0 END) AS queued,
                 SUM(CASE WHEN status = 'running' THEN 1 ELSE 0 END) AS running,
                 SUM(CASE WHEN status = 'success' THEN 1 ELSE 0 END) AS success,
                 SUM(CASE WHEN status = 'failed' THEN 1 ELSE 0 END) AS failed,
                 SUM(CASE WHEN status = 'canceled' THEN 1 ELSE 0 END) AS canceled,
                 COUNT(*) - SUM(
                   CASE WHEN status IN ('queued', 'running') THEN 1 ELSE 0 END
                 ) AS completed
          FROM platform_runs
          WHERE batch_id IS NOT NULL
          GROUP BY batch_id
        )
    """

    _RUN_BATCH_STATUS_EXPR = """
            CASE
              WHEN COALESCE(c.total, 0) = 0 THEN 'queued'
              WHEN COALESCE(c.queued, 0) = COALESCE(c.total, 0) THEN 'queued'
              WHEN COALESCE(c.running, 0) > 0
                OR (COALESCE(c.queued, 0) > 0 AND COALESCE(c.completed, 0) > 0)
                THEN 'running'
              WHEN COALESCE(c.success, 0) = COALESCE(c.total, 0) THEN 'success'
              WHEN COALESCE(c.canceled, 0) = COALESCE(c.total, 0) THEN 'canceled'
              WHEN COALESCE(c.completed, 0) = COALESCE(c.total, 0)
                AND COALESCE(c.success, 0) > 0 THEN 'partial_failed'
              WHEN COALESCE(c.completed, 0) = COALESCE(c.total, 0)
                AND COALESCE(c.failed, 0) > 0 THEN 'failed'
              ELSE 'running'
            END
    """

    def _run_batch_response(self, row: Any) -> dict[str, Any]:
        total = int(row[12] or 0)
        queued = int(row[13] or 0)
        running = int(row[14] or 0)
        success = int(row[15] or 0)
        failed = int(row[16] or 0)
        canceled = int(row[17] or 0)
        completed = success + failed + canceled
        return {
            "id": row[0],
            "projectId": row[1],
            "environmentId": row[2],
            "clientRequestId": row[3],
            "source": row[4],
            "retryOfBatchId": row[5],
            "flowIds": parse_json(row[6], []),
            "cancellationRequested": bool(row[7]),
            "createdBy": row[8],
            "createdAt": row[9],
            "updatedAt": row[10],
            "status": row[11],
            "counts": {
                "total": total,
                "queued": queued,
                "running": running,
                "success": success,
                "failed": failed,
                "canceled": canceled,
                "completed": completed,
            },
        }

    def _run_batch_select(self) -> str:
        status_expr = " ".join(
            line.strip() for line in self._RUN_BATCH_STATUS_EXPR.splitlines()
        ).strip()
        return f"""
        SELECT b.id, b.project_id, b.environment_id, b.client_request_id,
               b.source, b.retry_of_batch_id, b.requested_flow_ids,
               b.cancellation_requested, b.created_by, b.created_at, b.updated_at,
               {status_expr} AS status,
               COALESCE(c.total, 0) AS total, COALESCE(c.queued, 0) AS queued,
               COALESCE(c.running, 0) AS running, COALESCE(c.success, 0) AS success,
               COALESCE(c.failed, 0) AS failed, COALESCE(c.canceled, 0) AS canceled
        """

    def run_batch_by_id(self, project_id: str, batch_id: str) -> dict[str, Any]:
        from .http import PlatformError

        row = self.database.execute(
            f"""
            WITH {self._RUN_BATCH_COUNTS_CTE}
            {self._run_batch_select()}
            FROM run_batches b
            LEFT JOIN counts c ON c.batch_id = b.id
            WHERE b.id = ? AND b.project_id = ?
            """,
            (batch_id, project_id),
        ).fetchone()
        if not row:
            raise PlatformError(404, "RUN_BATCH_NOT_FOUND")
        return self._run_batch_response(row)

    def run_batches_page(
        self,
        project_id: str,
        page: int,
        page_size: int,
        status: str | None = None,
    ) -> dict[str, Any]:
        params: list[Any] = [project_id]
        status_filter = ""
        if status:
            status_filter = "WHERE status = ?"
            params.append(status)
        rows = self.database.execute(
            f"""
            WITH {self._RUN_BATCH_COUNTS_CTE}
            SELECT * FROM (
              {self._run_batch_select()}
              FROM run_batches b
              LEFT JOIN counts c ON c.batch_id = b.id
              WHERE b.project_id = ?
              ORDER BY b.created_at DESC, b.id DESC
            ) {status_filter}
            LIMIT ? OFFSET ?
            """,
            (*params, page_size, (page - 1) * page_size),
        ).fetchall()
        total_row = self.database.execute(
            f"""
            WITH {self._RUN_BATCH_COUNTS_CTE}
            SELECT COUNT(*) FROM (
              {self._run_batch_select()}
              FROM run_batches b
              LEFT JOIN counts c ON c.batch_id = b.id
              WHERE b.project_id = ?
            ) {status_filter}
            """,
            tuple(params),
        ).fetchone()
        return {
            "batches": [self._run_batch_response(row) for row in rows],
            "total": int(total_row[0] or 0),
            "page": page,
            "pageSize": page_size,
        }

    def batch_runs(self, project_id: str, batch_id: str) -> list[dict[str, Any]]:
        rows = self.database.execute(
            """
            SELECT id, project_id, revision_id, environment_id, agent_id,
                   executor_type, status, snapshot, cancellation_requested,
                   result, created_at, updated_at, batch_item_index,
                   retry_of_run_id
            FROM platform_runs
            WHERE batch_id = ? AND project_id = ?
            ORDER BY batch_item_index ASC
            """,
            (batch_id, project_id),
        ).fetchall()
        return [
            {
                "id": row[0],
                "projectId": row[1],
                "revisionId": row[2],
                "environmentId": row[3],
                "agentId": row[4],
                "executorType": row[5],
                "status": row[6],
                "snapshot": parse_json(row[7], {}),
                "cancellationRequested": bool(row[8]),
                "result": parse_json(row[9], None),
                "createdAt": row[10],
                "updatedAt": row[11],
                "batchItemIndex": row[12],
                "retryOfRunId": row[13],
            }
            for row in rows
        ]

    def _validate_batch_input(self, input: dict[str, Any]) -> list[str]:
        from .http import PlatformError

        flow_ids = input.get("flowIds")
        if not isinstance(flow_ids, list) or not all(
            isinstance(value, str) and value.strip() for value in flow_ids
        ):
            raise PlatformError(400, "BATCH_FLOW_IDS_INVALID")
        flow_ids = [value.strip() for value in flow_ids]
        if len(set(flow_ids)) != len(flow_ids):
            raise PlatformError(400, "BATCH_DUPLICATE_FLOW")
        if not (
            self.BATCH_MIN_FLOWS
            <= len(flow_ids)
            <= self.BATCH_MAX_FLOWS
        ):
            raise PlatformError(400, "BATCH_FLOW_COUNT_INVALID")
        if not (input.get("environmentId") or "").strip():
            raise PlatformError(400, "ENVIRONMENT_REQUIRED")
        if not (input.get("clientRequestId") or "").strip():
            raise PlatformError(400, "BATCH_CLIENT_REQUEST_ID_REQUIRED")
        if (
            input.get("revisionId")
            or input.get("datasetVersionId")
            or input.get("upToStepId")
        ):
            raise PlatformError(400, "BATCH_INPUT_NOT_SUPPORTED")
        return flow_ids

    def _existing_run_batch(
        self, project_id: str, client_request_id: str
    ) -> dict[str, Any] | None:
        existing = self.database.execute(
            """
            SELECT id FROM run_batches
            WHERE project_id = ? AND client_request_id = ?
            """,
            (project_id, client_request_id),
        ).fetchone()
        if not existing:
            return None
        return self.run_batch_by_id(project_id, existing[0])

    def _insert_run_batch(
        self,
        *,
        project_id: str,
        environment_id: str,
        client_request_id: str,
        flow_ids: list[str],
        created_by: str,
        specs: list[dict[str, Any]],
        retry_of_batch_id: str | None = None,
        retry_of_run_ids: list[str] | None = None,
        retry_rows: list[dict[str, Any]] | None = None,
    ) -> dict[str, Any]:
        import sqlite3

        from .http import PlatformError

        batch_id = str(uuid.uuid4())
        created_at = now()
        run_ids: list[str] = []
        self.database.execute("BEGIN IMMEDIATE")
        try:
            self.database.execute(
                """
                INSERT INTO run_batches (
                  id, project_id, environment_id, client_request_id, source,
                  retry_of_batch_id, requested_flow_ids, cancellation_requested,
                  created_by, created_at, updated_at
                ) VALUES (?, ?, ?, ?, 'manual', ?, ?, 0, ?, ?, ?)
                """,
                (
                    batch_id,
                    project_id,
                    environment_id,
                    client_request_id,
                    retry_of_batch_id,
                    json(flow_ids),
                    created_by,
                    created_at,
                    created_at,
                ),
            )
            for index, spec in enumerate(specs):
                run_ids.append(
                    self.insert_run_from_spec(
                        spec,
                        row=(
                            retry_rows[index]
                            if retry_rows
                            else {"rowNumber": None, "data": None}
                        ),
                        created_by=created_by,
                        source="manual",
                        batch_id=batch_id,
                        batch_item_index=index,
                        retry_of_run_id=(
                            retry_of_run_ids[index]
                            if retry_of_run_ids
                            else None
                        ),
                    )
                )
            self.database.execute("COMMIT")
        except sqlite3.IntegrityError:
            self.database.execute("ROLLBACK")
            existing = self._existing_run_batch(project_id, client_request_id)
            if existing is None:
                raise
            if (
                existing["environmentId"] != environment_id
                or existing["flowIds"] != flow_ids
            ):
                raise PlatformError(409, "IDEMPOTENCY_KEY_REUSED")
            return {
                "batch": existing,
                "runs": self.batch_runs(project_id, existing["id"]),
                "replayed": True,
            }
        except Exception:
            self.database.execute("ROLLBACK")
            raise
        for run_id in run_ids:
            self.enqueue_managed_run(run_id)
        return {
            "batch": self.run_batch_by_id(project_id, batch_id),
            "runs": self.batch_runs(project_id, batch_id),
            "replayed": False,
        }

    def create_run_batch(self, input: dict[str, Any]) -> dict[str, Any]:
        from .http import PlatformError

        project_id = input["projectId"]
        flow_ids = self._validate_batch_input(input)
        environment_id = input["environmentId"].strip()
        client_request_id = input["clientRequestId"].strip()
        existing = self._existing_run_batch(project_id, client_request_id)
        if existing is not None:
            if (
                existing["environmentId"] != environment_id
                or existing["flowIds"] != flow_ids
            ):
                raise PlatformError(409, "IDEMPOTENCY_KEY_REUSED")
            return {
                "batch": existing,
                "runs": self.batch_runs(project_id, existing["id"]),
                "replayed": True,
            }
        specs: list[dict[str, Any]] = []
        preflight_errors: list[dict[str, str]] = []
        total_steps = 0
        for flow_id in flow_ids:
            try:
                spec = self.resolve_run_spec(
                    {
                        "projectId": project_id,
                        "flowId": flow_id,
                        "environmentId": environment_id,
                    }
                )
                # A revision can carry a dataset default even when the batch
                # request does not. Batches are intentionally one run per flow.
                if spec["datasetVersionId"]:
                    raise PlatformError(400, "BATCH_INPUT_NOT_SUPPORTED")
                if not spec["flowSteps"]:
                    raise PlatformError(400, "FLOW_HAS_NO_STEPS")
                missing = self.missing_secret_names(
                    project_id,
                    [
                        name
                        for name in spec["secretNames"]
                        if name in spec["requiredSecretNames"]
                    ],
                )
                if missing:
                    raise PlatformError(409, "RUN_SECRET_NOT_CONFIGURED")
            except PlatformError as error:
                preflight_errors.append({"flowId": flow_id, "code": error.code})
                continue
            total_steps += len(spec["flowSteps"])
            specs.append(spec)
        if preflight_errors:
            raise PlatformError(
                409,
                "BATCH_PREFLIGHT_FAILED",
                {"items": preflight_errors},
            )
        if total_steps > self.BATCH_MAX_TOTAL_STEPS:
            raise PlatformError(413, "BATCH_TOTAL_STEPS_EXCEEDED")
        return self._insert_run_batch(
            project_id=project_id,
            environment_id=environment_id,
            client_request_id=client_request_id,
            flow_ids=flow_ids,
            created_by=input["createdBy"],
            specs=specs,
        )

    def cancel_run_batch(
        self, project_id: str, batch_id: str, actor_id: str
    ) -> dict[str, Any]:
        from .http import PlatformError

        self.run_batch_by_id(project_id, batch_id)
        self.database.execute("BEGIN IMMEDIATE")
        try:
            self.database.execute(
                """
                UPDATE run_batches
                SET cancellation_requested = 1, updated_at = ?
                WHERE id = ? AND cancellation_requested = 0
                """,
                (now(), batch_id),
            )
            queued_rows = self.database.execute(
                """
                SELECT id FROM platform_runs
                WHERE batch_id = ? AND project_id = ? AND status = 'queued'
                """,
                (batch_id, project_id),
            ).fetchall()
            if queued_rows:
                self.database.execute(
                    """
                    UPDATE platform_runs
                    SET status = 'canceled', updated_at = ?
                    WHERE batch_id = ? AND project_id = ?
                      AND status = 'queued'
                    """,
                    (now(), batch_id, project_id),
                )
            running_rows = self.database.execute(
                """
                SELECT id FROM platform_runs
                WHERE batch_id = ? AND project_id = ? AND status = 'running'
                  AND cancellation_requested = 0
                """,
                (batch_id, project_id),
            ).fetchall()
            if running_rows:
                self.database.execute(
                    """
                    UPDATE platform_runs
                    SET cancellation_requested = 1, updated_at = ?
                    WHERE batch_id = ? AND project_id = ? AND status = 'running'
                      AND cancellation_requested = 0
                    """,
                    (now(), batch_id, project_id),
                )
            affected = [row[0] for row in queued_rows] + [
                row[0] for row in running_rows
            ]
            for run_id in affected:
                self.append_run_event(
                    run_id, "run.cancel_requested", {"actorId": actor_id}
                )
            self.database.execute("COMMIT")
        except Exception:
            self.database.execute("ROLLBACK")
            raise
        for run_id in affected:
            # queued 子项也必须在提交后从 ManagedRunner 队列移除，
            # 否则 DB 已 canceled 但 runner 仍可能继续执行该流程。
            self.cancel_managed_run(run_id)
        return {
            "batch": self.run_batch_by_id(project_id, batch_id),
            "runs": self.batch_runs(project_id, batch_id),
            "affectedQueued": len(queued_rows),
            "affectedRunning": len(running_rows),
        }

    def retry_run_batch(
        self,
        project_id: str,
        batch_id: str,
        actor_id: str,
        client_request_id: str,
    ) -> dict[str, Any]:
        from .http import PlatformError

        batch = self.run_batch_by_id(project_id, batch_id)
        if batch["status"] not in ("success", "partial_failed", "failed", "canceled"):
            raise PlatformError(409, "BATCH_NOT_RETRYABLE")
        runs = self.batch_runs(project_id, batch_id)
        retry_items = [
            run for run in runs if run["status"] in ("failed", "canceled")
        ]
        if not retry_items:
            raise PlatformError(409, "BATCH_NOT_RETRYABLE")
        retry_flow_ids: list[str] = []
        for run in retry_items:
            flow_id = as_record(run["snapshot"].get("flow")).get("id")
            retry_flow_ids.append(
                flow_id if isinstance(flow_id, str) else run["revisionId"]
            )
        existing = self._existing_run_batch(project_id, client_request_id)
        if existing is not None:
            if (
                existing.get("retryOfBatchId") == batch_id
                and existing["flowIds"] == retry_flow_ids
            ):
                return {
                    "batch": existing,
                    "runs": self.batch_runs(project_id, existing["id"]),
                    "replayed": True,
                }
            raise PlatformError(409, "IDEMPOTENCY_KEY_REUSED")
        specs: list[dict[str, Any]] = []
        retry_rows: list[dict[str, Any]] = []
        for run in retry_items:
            spec = self.clone_retry_spec_from_run(run)
            self._preflight_retry_spec(project_id, spec)
            specs.append(spec)
            retry_rows.append(self._row_from_retry_snapshot(run["snapshot"]))
        return self._insert_run_batch(
            project_id=project_id,
            environment_id=batch["environmentId"],
            client_request_id=client_request_id,
            flow_ids=retry_flow_ids,
            created_by=actor_id,
            specs=specs,
            retry_of_batch_id=batch_id,
            retry_of_run_ids=[run["id"] for run in retry_items],
            retry_rows=retry_rows,
        )

    def run_by_id(
        self, run_id: str, project_id: str | None = None
    ) -> dict[str, Any]:
        from .http import PlatformError

        row = self.database.execute(
            """
            SELECT id, project_id, revision_id, environment_id, agent_id,
                   executor_type, status, snapshot, cancellation_requested,
                   result, created_at, updated_at, retry_of_run_id
            FROM platform_runs
            WHERE id = ? AND (? IS NULL OR project_id = ?)
            """,
            (run_id, project_id, project_id),
        ).fetchone()
        if not row:
            raise PlatformError(404, "RUN_NOT_FOUND")
        return {
            "id": row[0],
            "projectId": row[1],
            "revisionId": row[2],
            "environmentId": row[3],
            "agentId": row[4],
            "executorType": row[5],
            "status": row[6],
            "snapshot": parse_json(row[7], {}),
            "cancellationRequested": bool(row[8]),
            "result": parse_json(row[9], None),
            "createdAt": row[10],
            "updatedAt": row[11],
            "retryOfRunId": row[12],
        }


    def delete_run(self, project_id: str, run_id: str) -> dict[str, Any]:
        from .http import PlatformError

        run = self.database.execute(
            "SELECT id, project_id, status FROM platform_runs WHERE id = ?",
            (run_id,),
        ).fetchone()
        if not run:
            raise PlatformError(404, "RUN_NOT_FOUND")
        if run[1] != project_id:
            raise PlatformError(404, "RUN_NOT_FOUND")
        if run[2] not in _TERMINAL_RUN_STATUSES:
            raise PlatformError(409, "RUN_NOT_DELETABLE")
        with self.database:
            self.database.execute("DELETE FROM deliveries WHERE run_id = ?", (run_id,))
            self.database.execute("DELETE FROM flow_outputs WHERE run_id = ?", (run_id,))
            self.database.execute("DELETE FROM platform_artifacts WHERE run_id = ?", (run_id,))
            self.database.execute("DELETE FROM platform_run_events WHERE run_id = ?", (run_id,))
            self.database.execute("DELETE FROM platform_runs WHERE id = ? AND project_id = ?", (run_id, project_id))
        return {"runId": run_id, "deleted": True}

    def delete_runs(self, project_id: str, run_ids: list[str]) -> dict[str, Any]:
        if not run_ids:
            return {"runIds": [], "deletedCount": 0}
        placeholders = ",".join("?" for _ in run_ids)
        terminal_placeholders = ",".join("?" for _ in _TERMINAL_RUN_STATUSES)
        rows = self.database.execute(
            f"""
            SELECT id FROM platform_runs
            WHERE project_id = ?
              AND id IN ({placeholders})
              AND status IN ({terminal_placeholders})
            """,
            [project_id, *run_ids, *_TERMINAL_RUN_STATUSES],
        ).fetchall()
        deletable_run_ids = [row[0] for row in rows]
        if not deletable_run_ids:
            return {"runIds": [], "deletedCount": 0}
        deletable_placeholders = ",".join("?" for _ in deletable_run_ids)
        with self.database:
            self.database.execute(
                f"DELETE FROM deliveries WHERE run_id IN ({deletable_placeholders})",
                deletable_run_ids,
            )
            self.database.execute(
                f"DELETE FROM flow_outputs WHERE run_id IN ({deletable_placeholders})",
                deletable_run_ids,
            )
            self.database.execute(
                f"DELETE FROM platform_artifacts WHERE run_id IN ({deletable_placeholders})",
                deletable_run_ids,
            )
            self.database.execute(
                f"DELETE FROM platform_run_events WHERE run_id IN ({deletable_placeholders})",
                deletable_run_ids,
            )
            cursor = self.database.execute(
                f"DELETE FROM platform_runs WHERE id IN ({deletable_placeholders}) AND project_id = ?",
                [*deletable_run_ids, project_id],
            )
            deleted_count = cursor.rowcount
        return {"runIds": deletable_run_ids, "deletedCount": deleted_count}

    def run_response(self, run: dict[str, Any]) -> dict[str, Any]:
        agent = self.database.execute(
            """
            SELECT id, name, browser_version, os, max_concurrency, last_seen_at
            FROM agents WHERE id = ?
            """,
            (run["agentId"],),
        ).fetchone()
        artifacts = self.database.execute(
            """
            SELECT id, name, content_type, created_at FROM platform_artifacts
            WHERE run_id = ? ORDER BY created_at ASC
            """,
            (run["id"],),
        ).fetchall()
        events = self.database.execute(
            """
            SELECT id, kind, data, created_at FROM platform_run_events
            WHERE run_id = ? ORDER BY id ASC LIMIT 500
            """,
            (run["id"],),
        ).fetchall()
        flow_outputs = self.database.execute(
            """
            SELECT name, value, source, created_at FROM flow_outputs
            WHERE run_id = ? ORDER BY name ASC
            """,
            (run["id"],),
        ).fetchall()
        response: dict[str, Any] = {
            **run,
            "artifacts": [
                {
                    "id": row[0],
                    "name": row[1],
                    "contentType": row[2],
                    "createdAt": row[3],
                }
                for row in artifacts
            ],
            "events": [
                {
                    "id": row[0],
                    "kind": row[1],
                    "data": parse_json(row[2], {}),
                    "at": row[3],
                }
                for row in events
            ],
            "flowOutputs": [
                {
                    "name": row[0],
                    "value": row[1],
                    "source": row[2],
                    "createdAt": row[3],
                }
                for row in flow_outputs
            ],
        }
        if run["executorType"] == "agent" and agent:
            response["agent"] = {
                "id": agent[0],
                "name": agent[1],
                "browserVersion": agent[2],
                "os": agent[3],
                "maxConcurrency": agent[4],
                "lastSeenAt": agent[5],
            }
        return response

    def element_validation_by_id(
        self, validation_id: str, project_id: str | None = None
    ) -> dict[str, Any]:
        from .http import PlatformError

        row = self.database.execute(
            """
            SELECT id, project_id, environment_id, agent_id, status,
                   element_snapshot, result, error, created_at, updated_at
            FROM element_validations
            WHERE id = ? AND (? IS NULL OR project_id = ?)
            """,
            (validation_id, project_id, project_id),
        ).fetchone()
        if not row:
            raise PlatformError(404, "ELEMENT_VALIDATION_NOT_FOUND")
        return {
            "id": row[0],
            "projectId": row[1],
            "environmentId": row[2],
            "agentId": row[3],
            "status": row[4],
            "element": parse_json(row[5], {}),
            "result": parse_json(row[6], None),
            "error": row[7] or None,
            "createdAt": row[8],
            "updatedAt": row[9],
        }

    def create_element_validation(
        self, project_id: str, environment_id: str, element: dict[str, Any], created_by: str
    ) -> dict[str, Any]:
        from .http import PlatformError

        row = self.database.execute(
            """
            SELECT data FROM project_resources
            WHERE project_id = ? AND resource_type = 'environments'
              AND resource_id = ? AND archived_at IS NULL
            """,
            (project_id, environment_id),
        ).fetchone()
        if row:
            environment = parse_json(row[0], {})
        else:
            document = self.document_for(project_id)
            environments = document["data"].get("environments", [])
            environment = next(
                (
                    item
                    for item in environments
                    if isinstance(item, dict) and item.get("id") == environment_id
                ),
                None,
            )
        if not isinstance(environment, dict):
            raise PlatformError(404, "ENVIRONMENT_NOT_FOUND")
        self.require_chromium_environment(environment)
        self.require_same_origin_element_path(environment, element)
        agent = self.managed_agent(project_id)
        validation_id = str(uuid.uuid4())
        created_at = now()
        self.database.execute(
            """
            INSERT INTO element_validations (
              id, project_id, environment_id, agent_id, status,
              element_snapshot, created_by, created_at, updated_at
            ) VALUES (?, ?, ?, ?, 'queued', ?, ?, ?, ?)
            """,
            (
                validation_id,
                project_id,
                environment_id,
                agent["id"],
                json(element),
                created_by,
                created_at,
                created_at,
            ),
        )
        validation = self.element_validation_by_id(validation_id)
        self.enqueue_managed_validation(validation, environment)
        return self.element_validation_by_id(validation_id)

    def require_same_origin_element_path(
        self,
        environment: dict[str, Any],
        element: dict[str, Any],
    ) -> None:
        from .http import PlatformError

        base_url = str(environment.get("baseUrl", ""))
        element_path = str(element.get("path", "/"))
        try:
            base = urlsplit(base_url)
            target = urlsplit(urljoin(base_url, element_path))
            if (
                base.scheme not in ("http", "https")
                or target.scheme != base.scheme
                or target.netloc != base.netloc
            ):
                raise PlatformError(400, "ELEMENT_VALIDATION_TARGET_FORBIDDEN")
        except PlatformError:
            raise
        except Exception:
            raise PlatformError(400, "ELEMENT_VALIDATION_TARGET_INVALID") from None

    def cancel_managed_run(self, run_id: str) -> bool:
        return self.managed_runner.cancel(run_id)

    def managed_runner_input(self, run: dict[str, Any]) -> dict[str, Any]:
        snapshot = run["snapshot"]
        flow = as_record(snapshot.get("flow"))
        environment = as_record(snapshot.get("environment"))
        variable_rows = self.database.execute(
            """
            SELECT data FROM project_resources
            WHERE project_id = ? AND resource_type = 'variables'
              AND archived_at IS NULL
            """,
            (run["projectId"],),
        ).fetchall()
        variables: dict[str, str] = {}
        for variable_row in variable_rows:
            variable = parse_json(variable_row[0], {})
            if (
                not isinstance(variable, dict)
                or variable.get("secret") is True
                or not isinstance(variable.get("name"), str)
                or not isinstance(variable.get("value"), str)
            ):
                continue
            scope = (
                "env"
                if variable.get("scope") == "环境"
                else "project"
                if variable.get("scope") == "项目"
                else ""
            )
            key = f"{scope}.{variable['name']}" if scope else variable["name"]
            variables[key] = variable["value"]
        secret_names = snapshot.get("secretNames", [])
        if not isinstance(secret_names, list):
            secret_names = []
        secret_names = [name for name in secret_names if isinstance(name, str)]
        dataset_row = as_record(as_record(snapshot.get("datasetRow")).get("data"))
        return {
            "environment": environment,
            "flow": {
                "id": flow.get("id") if isinstance(flow.get("id"), str) else run["revisionId"],
                "name": flow.get("name") if isinstance(flow.get("name"), str) else "Published flow",
                "steps": flow.get("steps") if isinstance(flow.get("steps"), list) else [],
            },
            "elements": snapshot.get("elements")
            if isinstance(snapshot.get("elements"), list)
            else [],
            "variables": variables,
            "data": {
                str(key): str(value or "")
                for key, value in dataset_row.items()
            },
            "secrets": self.secret_values(run["projectId"], secret_names),
            "upToStepId": (
                snapshot.get("upToStepId")
                if isinstance(snapshot.get("upToStepId"), str)
                else None
            ),
        }

    def enqueue_managed_run(self, run_id: str) -> None:
        run = self.run_by_id(run_id)
        if run["status"] != "queued":
            return
        input = self.managed_runner_input(run)

        def started() -> None:
            self.database.execute(
                """
                UPDATE platform_runs SET status = 'running', updated_at = ?
                WHERE id = ? AND status = 'queued'
                """,
                (now(), run_id),
            )
            self.append_run_event(run_id, "run.started", {"executorType": "managed"})

        def event(kind: str, data: dict[str, Any]) -> None:
            self.append_run_event(
                run_id,
                kind,
                self.redact_run_value(run, data),
            )

        def artifact(input_data: dict[str, Any]) -> None:
            artifact_id = str(uuid.uuid4())
            self.database.execute(
                """
                INSERT INTO platform_artifacts (
                  id, run_id, project_id, name, content_type, path, created_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    artifact_id,
                    run_id,
                    run["projectId"],
                    safe_artifact_name(str(input_data["name"])),
                    str(input_data["contentType"])[:120],
                    str(input_data["path"]),
                    now(),
                ),
            )
            self.append_run_event(
                run_id,
                "artifact.created",
                {
                    "artifactId": artifact_id,
                    "name": str(input_data["name"]),
                    "contentType": str(input_data["contentType"]),
                },
            )

        def completed(result: dict[str, Any]) -> None:
            current_run = self.run_by_id(run_id)
            safe_result = self.redact_run_value(current_run, result)
            requested_status = (
                result.get("status")
                if result.get("status") in ("success", "failed")
                else "failed"
            )
            status = "canceled" if current_run["cancellationRequested"] else requested_status
            updated = self.database.execute(
                """
                UPDATE platform_runs
                SET status = ?, result = ?, updated_at = ?
                WHERE id = ? AND status IN ('queued', 'running')
                """,
                (status, json(safe_result), now(), run_id),
            )
            if updated.rowcount != 1:
                return
            self.persist_flow_outputs(current_run, safe_result)
            self.append_run_event(
                run_id,
                "run.complete",
                {"status": status, "result": safe_result, "executorType": "managed"},
            )
            self.audit_run_lifecycle(run_id, current_run, status)
            self.queue_run_deliveries(self.run_by_id(run_id), status)

        self.managed_runner.enqueue(
            run_id,
            input,
            {
                "started": started,
                "event": event,
                "artifact": artifact,
                "completed": completed,
            },
            kind="run",
            workspace_id=self.project_for(run["projectId"])["workspace_id"],
        )

    def enqueue_managed_validation(
        self,
        validation: dict[str, Any],
        environment: dict[str, Any],
    ) -> None:
        validation_id = validation["id"]

        def started() -> None:
            self.database.execute(
                """
                UPDATE element_validations SET status = 'running', updated_at = ?
                WHERE id = ? AND status = 'queued'
                """,
                (now(), validation_id),
            )

        def artifact(input_data: dict[str, Any]) -> None:
            self.database.execute(
                """
                INSERT INTO element_validation_artifacts (
                  id, validation_id, project_id, name, content_type, path,
                  created_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?)
                """,
                (
                    str(uuid.uuid4()),
                    validation_id,
                    validation["projectId"],
                    safe_artifact_name(str(input_data["name"])),
                    str(input_data["contentType"])[:120],
                    str(input_data["path"]),
                    now(),
                ),
            )

        def completed(result: dict[str, Any]) -> None:
            artifact_row = self.database.execute(
                """
                SELECT id FROM element_validation_artifacts
                WHERE validation_id = ? ORDER BY created_at DESC LIMIT 1
                """,
                (validation_id,),
            ).fetchone()
            payload = {
                "count": result.get("count"),
                "firstMatch": result.get("firstMatch"),
                "elapsedMs": result.get("elapsedMs"),
                "screenshotId": artifact_row[0] if artifact_row else None,
            }
            self.database.execute(
                """
                UPDATE element_validations
                SET status = ?, result = ?, error = ?, updated_at = ?
                WHERE id = ?
                """,
                (
                    result.get("status"),
                    json(payload),
                    result.get("error"),
                    now(),
                    validation_id,
                ),
            )

        self.managed_runner.enqueue(
            validation_id,
            {"environment": environment, "element": validation["element"]},
            {
                "started": started,
                "artifact": artifact,
                "event": lambda *_args: None,
                "completed": completed,
            },
            kind="validation",
            workspace_id=self.project_for(validation["projectId"])["workspace_id"],
        )

    def metrics(self) -> dict[str, Any]:
        """OBS-02 service-level metrics (DB-backed, JSON)."""
        run_counts: dict[str, int] = {}
        for row in self.database.execute(
            "SELECT status, COUNT(*) FROM platform_runs GROUP BY status"
        ).fetchall():
            run_counts[str(row[0])] = int(row[1])

        delivery_counts: dict[str, int] = {}
        for row in self.database.execute(
            "SELECT status, COUNT(*) FROM deliveries GROUP BY status"
        ).fetchall():
            delivery_counts[str(row[0])] = int(row[1])

        disk: dict[str, int] | None = None
        try:
            usage = shutil.disk_usage(str(self.data_directory))
            disk = {"total": usage.total, "used": usage.used, "free": usage.free}
        except Exception:
            pass

        return {
            "runs": run_counts,
            "deliveries": delivery_counts,
            "disk": disk,
            "artifactBytes": self._artifact_bytes(),
        }

    def _artifact_bytes(self) -> int:
        total = 0
        try:
            for path in self.managed_runner.artifact_directory.rglob("*"):
                if path.is_file():
                    total += path.stat().st_size
        except Exception:
            pass
        return total

    def redact_run_value(self, run: dict[str, Any], value: Any) -> Any:
        try:
            rows = self.database.execute(
                """
                SELECT name, iv, tag, ciphertext FROM project_secrets
                WHERE project_id = ?
                """,
                (run["projectId"],),
            ).fetchall()
            secrets = {
                row[0]: self.decrypt({"iv": row[1], "tag": row[2], "ciphertext": row[3]})
                for row in rows
            }

            def redact(current: Any) -> Any:
                if isinstance(current, str):
                    result = current
                    for secret in secrets.values():
                        if secret:
                            result = result.replace(secret, "***")
                    return result
                if isinstance(current, list):
                    return [redact(item) for item in current]
                if isinstance(current, dict):
                    return {key: redact(item) for key, item in current.items()}
                return current

            return redact(value)
        except Exception:
            return "***"

    def retention_cleanup(
        self,
        audit_days: int = 180,
        run_days: int = 90,
        artifact_days: int = 15,
        dry_run: bool = False,
    ) -> dict[str, int]:
        """DATA-01 retention pass.

        Removes expired artifacts (file + row), runs (cascade events/outputs/
        artifacts), audit events, expired sessions and delivered notifications.
        Returns counts; when ``dry_run`` is True nothing is deleted.
        """
        artifact_directory = self.managed_runner.artifact_directory

        def cutoff(days: int) -> str:
            return (
                datetime.now(timezone.utc) - timedelta(days=days)
            ).strftime("%Y-%m-%dT%H:%M:%S.%f")[:-3] + "Z"

        artifact_cutoff = cutoff(artifact_days)
        run_cutoff = cutoff(run_days)
        audit_cutoff = cutoff(audit_days)

        summary = {
            "artifacts": 0,
            "runs": 0,
            "runEvents": 0,
            "flowOutputs": 0,
            "auditEvents": 0,
            "sessions": 0,
            "deliveries": 0,
        }

        def _delete_artifact_file(path_value: str | None) -> None:
            if not path_value:
                return
            try:
                candidate = Path(path_value).resolve()
                base = Path(artifact_directory).resolve()
                if str(candidate).startswith(str(base) + os.sep) and candidate.is_file():
                    candidate.unlink()
            except Exception:
                pass

        # 1. Expired artifacts (shortest retention) — file + row.
        artifact_rows = self.database.execute(
            "SELECT id, path FROM platform_artifacts WHERE created_at <= ?",
            (artifact_cutoff,),
        ).fetchall()
        for row in artifact_rows:
            summary["artifacts"] += 1
            if not dry_run:
                _delete_artifact_file(row[1])
                self.database.execute(
                    "DELETE FROM platform_artifacts WHERE id = ?", (row[0],)
                )

        # 2. Expired runs — cascade events/outputs/artifacts, then the run.
        run_ids = self.database.execute(
            "SELECT id FROM platform_runs WHERE created_at <= ?",
            (run_cutoff,),
        ).fetchall()
        for run_row in run_ids:
            run_id = run_row[0]
            summary["runs"] += 1
            if dry_run:
                continue
            run_artifact_rows = self.database.execute(
                "SELECT id, path FROM platform_artifacts WHERE run_id = ?",
                (run_id,),
            ).fetchall()
            for artifact_row in run_artifact_rows:
                _delete_artifact_file(artifact_row[1])
                self.database.execute(
                    "DELETE FROM platform_artifacts WHERE id = ?", (artifact_row[0],)
                )
            summary["runEvents"] += self.database.execute(
                "DELETE FROM platform_run_events WHERE run_id = ?", (run_id,)
            ).rowcount
            summary["flowOutputs"] += self.database.execute(
                "DELETE FROM flow_outputs WHERE run_id = ?", (run_id,)
            ).rowcount
            self.database.execute(
                "DELETE FROM deliveries WHERE run_id = ?", (run_id,)
            )
            self.database.execute(
                "DELETE FROM platform_runs WHERE id = ?", (run_id,)
            )

        # 3. Expired audit events.
        summary["auditEvents"] = self.database.execute(
            "DELETE FROM audit_events WHERE created_at <= ?", (audit_cutoff,)
        ).rowcount

        # 4. Expired sessions and delivered notifications.
        summary["sessions"] = self.database.execute(
            "DELETE FROM platform_sessions WHERE expires_at <= ?", (now(),)
        ).rowcount
        summary["deliveries"] = self.database.execute(
            """
            DELETE FROM deliveries
            WHERE status IN ('delivered', 'failed') AND created_at <= ?
            """,
            (run_cutoff,),
        ).rowcount

        return summary

    def persist_flow_outputs(
        self,
        run: dict[str, Any],
        result: dict[str, Any],
    ) -> None:
        outputs = as_record(result.get("flowOutputs"))
        allowed_names = public_flow_output_names(run)
        for name, source_value in outputs.items():
            output_name = str(name).strip()[:120]
            if output_name not in allowed_names:
                continue
            if isinstance(source_value, str):
                value = source_value
            elif isinstance(source_value, (int, float, bool)):
                value = str(source_value)
            else:
                value = ""
            value = self.redact_run_value(run, value)
            if not value:
                continue
            self.database.execute(
                """
                INSERT INTO flow_outputs (
                  id, run_id, name, value, source, created_at
                ) VALUES (?, ?, ?, ?, 'agent', ?)
                ON CONFLICT(run_id, name) DO UPDATE SET
                  value = excluded.value, source = excluded.source,
                  created_at = excluded.created_at
                """,
                (str(uuid.uuid4()), run["id"], output_name, str(value)[:20000], now()),
            )

    def notification_payload(
        self,
        run: dict[str, Any],
        status: str,
    ) -> dict[str, Any]:
        latest_failure = self.database.execute(
            """
            SELECT kind, data FROM platform_run_events
            WHERE run_id = ? AND (kind LIKE '%failed%' OR kind LIKE '%error%')
            ORDER BY id DESC LIMIT 1
            """,
            (run["id"],),
        ).fetchone()
        artifacts = self.database.execute(
            """
            SELECT id, name FROM platform_artifacts
            WHERE run_id = ? ORDER BY created_at ASC
            """,
            (run["id"],),
        ).fetchall()
        return {
            "runId": run["id"],
            "projectId": run["projectId"],
            "status": status,
            "environmentId": run["environmentId"],
            "revisionId": run["revisionId"],
            "agentId": run["agentId"],
            "failedStep": (
                {
                    "kind": latest_failure[0],
                    "data": self.redact_run_value(
                        run, parse_json(latest_failure[1], {})
                    ),
                }
                if latest_failure
                else None
            ),
            "artifacts": [
                {"id": row[0], "name": row[1]} for row in artifacts
            ],
            "retry": {"cancellationRequested": run["cancellationRequested"]},
            "completedAt": now(),
        }

    def queue_run_deliveries(self, run: dict[str, Any], status: str) -> None:
        if status not in ("success", "failed", "canceled"):
            return
        project = self.project_for(run["projectId"])
        subscriptions = self.database.execute(
            """
            SELECT c.id FROM notification_subscriptions s
            JOIN notification_channels c ON c.id = s.channel_id
            WHERE s.project_id = ? AND c.workspace_id = ? AND c.enabled = 1
              AND c.archived_at IS NULL
              AND ((? = 'success' AND s.on_success = 1)
                   OR (? != 'success' AND s.on_failure = 1))
            """,
            (run["projectId"], project["workspace_id"], status, status),
        ).fetchall()
        payload = json(self.notification_payload(run, status))
        for subscription in subscriptions:
            self.database.execute(
                """
                INSERT INTO deliveries (
                  id, channel_id, run_id, status, payload, next_attempt_at,
                  created_at, updated_at
                )
                SELECT ?, ?, ?, 'pending', ?, ?, ?, ?
                WHERE NOT EXISTS (
                  SELECT 1 FROM deliveries WHERE channel_id = ? AND run_id = ?
                )
                """,
                (
                    str(uuid.uuid4()),
                    subscription[0],
                    run["id"],
                    payload,
                    now(),
                    now(),
                    now(),
                    subscription[0],
                    run["id"],
                ),
            )
        self.deliver_pending_notifications()

    def audit_run_lifecycle(
        self,
        run_id: str,
        run: dict[str, Any],
        status: str,
    ) -> None:
        project = self.project_for(run["projectId"])
        detail: dict[str, Any] = {"status": status}
        if status == "failed":
            failure = self.database.execute(
                """
                SELECT kind, data FROM platform_run_events
                WHERE run_id = ? AND (kind LIKE '%failed%' OR kind LIKE '%error%')
                ORDER BY id DESC LIMIT 1
                """,
                (run_id,),
            ).fetchone()
            failure_data = parse_json(failure[1], {}) if failure else {}
            if isinstance(failure_data, dict):
                if isinstance(failure_data.get("code"), str) and failure_data["code"]:
                    detail["errorCode"] = failure_data["code"]
                elif isinstance(failure_data.get("reason"), str) and failure_data["reason"]:
                    detail["errorCode"] = failure_data["reason"][:200]
                if isinstance(failure_data.get("stepId"), str) and failure_data["stepId"]:
                    detail["stepId"] = failure_data["stepId"]
        self.audit(
            project["workspace_id"],
            {"type": "system", "id": "managed-runner"},
            f"run.{'canceled' if status == 'canceled' else 'failed' if status == 'failed' else 'completed'}",
            {"type": "run_batch", "id": run_id},
            detail,
            run["projectId"],
        )

    def finalize_run_as_interrupted(self, run_id: str, reason: str) -> None:
        updated = self.database.execute(
            """
            UPDATE platform_runs
            SET status = 'failed', result = ?, updated_at = ?
            WHERE id = ? AND status IN ('queued', 'running')
            """,
            (json({"error": reason, "interrupted": True}), now(), run_id),
        )
        if updated.rowcount != 1:
            return
        self.append_run_event(run_id, "run.interrupted", {"reason": reason})
        self.append_run_event(
            run_id, "run.failed", {"reason": reason, "interrupted": True}
        )
        self.audit_run_lifecycle(run_id, self.run_by_id(run_id), "failed")
        self.queue_run_deliveries(self.run_by_id(run_id), "failed")

    def parse_dataset_upload(
        self, file_name: str, content_base64: str
    ) -> dict[str, Any]:
        from .http import PlatformError

        content = base64.b64decode(content_base64)
        if not content:
            raise PlatformError(400, "DATASET_FILE_EMPTY")
        if len(content) > 12_000_000:
            raise PlatformError(413, "DATASET_FILE_TOO_LARGE")
        match = re.search(r"\.([a-z0-9]+)$", file_name.lower())
        extension = match.group(1) if match else ""
        if extension not in ("csv", "xlsx"):
            raise PlatformError(400, "DATASET_FILE_TYPE_UNSUPPORTED")
        try:
            if extension == "csv":
                rows = parse_csv(content.decode("utf-8"))
            else:
                rows = _read_xlsx(content)
        except Exception as exc:
            if isinstance(exc, PlatformError):
                raise
            raise PlatformError(400, "DATASET_FILE_INVALID") from exc
        return {
            **normalize_dataset_rows(rows),
            "sourceName": safe_artifact_name(file_name),
        }

    @staticmethod
    def require_chromium_environment(environment: dict[str, Any]) -> None:
        from .http import PlatformError

        browser = environment.get("browser")
        if not isinstance(browser, str):
            browser = "Chromium"
        if browser != "Chromium":
            raise PlatformError(400, "AGENT_BROWSER_UNSUPPORTED")

    def ensure_chromium_available(self) -> None:
        import platform as _platform
        import sys as _sys
        global _CHROMIUM_AVAILABLE, _CHROMIUM_AVAILABLE_AT
        from .http import PlatformError

        now = _time.monotonic()
        cached_value = _CHROMIUM_AVAILABLE
        cached_at = _CHROMIUM_AVAILABLE_AT
        if cached_value is not None and cached_at is not None:
            ttl = (
                _CHROMIUM_CACHE_TTL_POSITIVE
                if cached_value
                else _CHROMIUM_CACHE_TTL_NEGATIVE
            )
            if now - cached_at < ttl:
                if not cached_value:
                    raise PlatformError(409, "AGENT_BROWSER_UNSUPPORTED")
                return
        available: bool = False
        executable: Path | None = None
        probe_error: Exception | None = None
        system_name: str = _platform.system()
        browsers_root: Path = Path(
            os.environ.get("PLAYWRIGHT_BROWSERS_PATH")
            or (Path.home() / ".cache" / "ms-playwright")
        )
        # 根据平台选择 Chromium 二进制相对路径。若添加新浏览器类型，同步更新。
        if system_name == "Linux":
            candidate_rel = Path("chrome-linux64") / "chrome"
        elif system_name == "Darwin":
            candidate_rel = (
                Path("chrome-mac") / "Chromium.app" / "Contents" / "MacOS" / "Chromium"
            )
        elif system_name == "Windows":
            candidate_rel = Path("chrome-win64") / "chrome.exe"
        else:
            candidate_rel = Path("chrome")  # 占位，下面会抛出 unsupported platform
        try:
            # 注意：在 uvicorn async handler 中同步调用 sync_playwright() 会被
            # Playwright 直接拒绝（"It looks like you are using Playwright Sync API
            # inside the asyncio loop."），因此这里不启动 Playwright 运行时，
            # 而是直接解析 Playwright 官方浏览器缓存目录结构定位二进制。
            # 缓存位置 & 二进制相对路径已在函数外层按平台计算好（browsers_root /
            # candidate_rel）。如果不支持的平台，在这里直接抛错走 fallback。
            if system_name not in {"Linux", "Darwin", "Windows"}:
                raise RuntimeError(f"Unsupported platform system={system_name!r}")
            # Playwright 的子目录命名是 chromium-<revision_number>，取所有匹配项中
            # 路径存在（按 mtime 取最新）的那一个。如果没有设置 PLAYWRIGHT_BROWSERS_PATH
            # 且 ms-playwright 也不存在，可能用户用 PLAYWRIGHT_BROWSERS_PATH=/nonexistent
            # 覆盖了默认，此时 fallback 到 async API 以兼容自定义安装。
            found_candidates: list[Path] = []
            if browsers_root.exists():
                for entry in browsers_root.iterdir():
                    if entry.is_dir() and entry.name.startswith("chromium-"):
                        candidate = entry / candidate_rel
                        if candidate.exists():
                            found_candidates.append(candidate)
            if found_candidates:
                found_candidates.sort(key=lambda p: p.stat().st_mtime, reverse=True)
                executable = found_candidates[0]
                available = True
            else:
                # Fallback：如果目录扫描没找到（例如非标准的 PLAYWRIGHT_BROWSERS_PATH
                # 映射、或自定义编译的 Chromium），则在新线程中通过 asyncio event loop
                # 调用 Playwright async API 查 executable_path — 这样不会与主线程
                # uvicorn 已运行的 event loop 冲突。
                try:
                    import asyncio as _asyncio
                    import threading as _threading

                    async def _probe_async() -> str | None:
                        from playwright.async_api import async_playwright

                        async with async_playwright() as pw:
                            return pw.chromium.executable_path  # type: ignore[no-any-return]

                    def _runner(loop: _asyncio.AbstractEventLoop) -> str | None:
                        future = _asyncio.run_coroutine_threadsafe(_probe_async(), loop)
                        try:
                            return future.result(timeout=30)
                        finally:
                            loop.call_soon_threadsafe(loop.stop)

                    new_loop = _asyncio.new_event_loop()
                    thread = _threading.Thread(
                        target=new_loop.run_forever, daemon=True
                    )
                    thread.start()
                    try:
                        async_path = _runner(new_loop)
                    finally:
                        thread.join(timeout=5)
                    if async_path:
                        executable = Path(async_path)
                        available = executable.exists()
                except Exception as exc:  # noqa: BLE001
                    probe_error = exc
                    available = False
        except Exception as exc:  # noqa: BLE001
            probe_error = exc
            available = False
        _CHROMIUM_AVAILABLE = available
        _CHROMIUM_AVAILABLE_AT = now
        if not available:
            details: list[str] = [f"[autoflow:chromium] probe failed at {now:.1f}s (pid {os.getpid()})"]
            if probe_error is not None:
                details.append(f"  exception: {type(probe_error).__name__}: {probe_error}")
            elif executable is None:
                details.append("  no chromium-<rev> directory found in browsers root")
                details.append(f"  browsers_root={browsers_root!s} (exists={browsers_root.exists()})")
                details.append(f"  expected relative binary: {candidate_rel}")
            else:
                details.append(f"  executable not found on disk: {executable}")
                details.append(f"  file exists={executable.exists()}, parent exists={executable.parent.exists()}")
            details.append(f"  system={system_name} HOME={os.environ.get('HOME')!r} PLAYWRIGHT_BROWSERS_PATH={os.environ.get('PLAYWRIGHT_BROWSERS_PATH')!r}")
            _sys.stderr.write("\n".join(details) + "\n")
            _sys.stderr.flush()
            raise PlatformError(409, "AGENT_BROWSER_UNSUPPORTED")

    def put_document(
        self,
        project_id: str,
        data: dict[str, Any],
        expected_version: int | None = None,
    ) -> dict[str, Any]:
        from .http import PlatformError

        current = self.document_for(project_id)
        if expected_version is not None and expected_version != current["version"]:
            raise PlatformError(409, "DOCUMENT_VERSION_CONFLICT")
        version = current["version"] + 1
        self.database.execute(
            """
            INSERT INTO project_documents (project_id, data, version, updated_at)
            VALUES (?, ?, ?, ?)
            ON CONFLICT(project_id) DO UPDATE SET
              data = excluded.data,
              version = excluded.version,
              updated_at = excluded.updated_at
            """,
            (project_id, json(data), version, now()),
        )
        migrate_project_document_resources(self.database, project_id, data)
        return {"version": version, "data": data}

    def project_analytics(
        self,
        project_id: str,
        options: dict[str, Any] | None = None,
    ) -> dict[str, Any]:
        options = options or {}
        period = "week" if options.get("period") == "week" else "day"
        category_by = options.get("categoryBy")
        if category_by not in ("code", "step"):
            category_by = "message"
        window_days = options.get("windowDays")
        from_iso = options.get("from")
        if from_iso is None and window_days:
            from_iso = _iso_add_seconds(-int(window_days) * 86400)
        to_iso = options.get("to")
        limit = options.get("limit")
        if limit is not None:
            limit = min(2000, max(1, int(limit)))
        runs = _fetch_runs(
            self.database, project_id, from_iso, to_iso, limit or 2000
        )
        events_by_run: dict[str, list[dict[str, Any]]] = {}
        for run in runs:
            rows = self.database.execute(
                """
                SELECT kind, data, created_at FROM platform_run_events
                WHERE run_id = ? ORDER BY id ASC
                """,
                (run["id"],),
            ).fetchall()
            events_by_run[run["id"]] = [
                {
                    "kind": row[0],
                    "data": parse_json(row[1], {}),
                    "at": row[2],
                }
                for row in rows
            ]

        trend: dict[str, dict[str, Any]] = {}
        categories: dict[str, int] = {}
        slow_steps: dict[str, dict[str, Any]] = {}
        elements: dict[str, dict[str, Any]] = {}
        durations: dict[str, dict[str, Any]] = {}
        for run in runs:
            date = _period_key(run["created_at"], period)
            point = trend.setdefault(
                date,
                {"date": date, "total": 0, "success": 0, "failed": 0, "canceled": 0},
            )
            point["total"] += 1
            if run["status"] == "success":
                point["success"] += 1
            elif run["status"] == "failed":
                point["failed"] += 1
            elif run["status"] == "canceled":
                point["canceled"] += 1

            events = events_by_run.get(run["id"], [])
            first = events[0] if events else None
            terminal = next(
                (
                    event
                    for event in reversed(events)
                    if event["kind"]
                    in ("run.complete", "run.failed", "run.interrupted")
                ),
                None,
            )
            if first and terminal:
                start_ms = _iso_ms(first["at"])
                end_ms = _iso_ms(terminal["at"])
                if start_ms is not None and end_ms is not None and end_ms >= start_ms:
                    duration_point = durations.setdefault(
                        date, {"date": date, "totalMs": 0, "count": 0}
                    )
                    duration_point["totalMs"] += end_ms - start_ms
                    duration_point["count"] += 1

            failure = next(
                (
                    event
                    for event in reversed(events)
                    if event["kind"] == "run.failed" or "error" in event["kind"]
                ),
                None,
            )
            if failure:
                data = failure["data"]
                if category_by == "step":
                    category = str(data.get("stepId") or "unknown")
                elif category_by == "code":
                    category = failure_category(
                        data.get("message"), data.get("code")
                    )
                else:
                    category = failure_category(data.get("message"))
                categories[category] = categories.get(category, 0) + 1

            for event in events:
                if event["kind"] not in ("step.completed", "step.succeeded"):
                    continue
                duration_ms = event["data"].get("durationMs")
                try:
                    duration_ms = float(duration_ms)
                except (TypeError, ValueError):
                    continue
                if duration_ms < 0:
                    continue
                step_id = str(event["data"].get("stepId") or "unknown")
                current = slow_steps.setdefault(
                    step_id,
                    {
                        "stepId": step_id,
                        "title": str(event["data"].get("title") or step_id),
                        "totalMs": 0,
                        "maxMs": 0,
                        "count": 0,
                    },
                )
                current["totalMs"] += duration_ms
                current["maxMs"] = max(current["maxMs"], duration_ms)
                current["count"] += 1

            snapshot = parse_json(run["snapshot"], {})
            flow = snapshot.get("flow") if isinstance(snapshot, dict) else {}
            if not isinstance(flow, dict):
                flow = {}
            step_definitions = flow.get("steps", [])
            if not isinstance(step_definitions, list):
                step_definitions = []
            element_definitions = snapshot.get("elements", [])
            if not isinstance(element_definitions, list):
                element_definitions = []
            failed_step_ids = {
                str(event["data"].get("stepId") or "")
                for event in events
                if event["kind"] == "run.failed"
            }
            for raw_step in step_definitions:
                step = raw_step if isinstance(raw_step, dict) else {}
                reference = step.get("element")
                if not isinstance(reference, str):
                    reference = step.get("elementId")
                if not isinstance(reference, str) or not reference:
                    continue
                definition = next(
                    (
                        item
                        for item in element_definitions
                        if isinstance(item, dict)
                        and (
                            item.get("id") == reference
                            or item.get("name") == reference
                        )
                    ),
                    None,
                )
                if isinstance(definition, dict):
                    element_id = (
                        definition.get("id")
                        if isinstance(definition.get("id"), str)
                        else reference
                    )
                    name = (
                        definition.get("name")
                        if isinstance(definition.get("name"), str)
                        else reference
                    )
                else:
                    element_id = reference
                    name = reference
                current_element = elements.setdefault(
                    element_id,
                    {
                        "elementId": element_id,
                        "name": name,
                        "runCount": 0,
                        "flowCount": set(),
                        "failedRuns": 0,
                        "lastUsedAt": run["created_at"],
                    },
                )
                current_element["runCount"] += 1
                current_element["flowCount"].add(run["revision_id"])
                if str(step.get("id") or "") in failed_step_ids:
                    current_element["failedRuns"] += 1
                if run["created_at"] > current_element["lastUsedAt"]:
                    current_element["lastUsedAt"] = run["created_at"]

        previous = None
        from_ms = _iso_ms(from_iso) if from_iso else None
        to_ms = _iso_ms(to_iso) if to_iso else None
        if (
            from_ms is not None
            and to_ms is not None
            and to_ms > from_ms
        ):
            previous = _summarize_runs(
                _fetch_runs(
                    self.database,
                    project_id,
                    _iso_from_ms(from_ms - (to_ms - from_ms)),
                    from_iso,
                )
            )
        elif from_ms is not None and options.get("windowDays"):
            previous = _summarize_runs(
                _fetch_runs(
                    self.database,
                    project_id,
                    _iso_from_ms(from_ms - int(options["windowDays"]) * 86400000),
                    from_iso,
                )
            )
        elif limit and runs:
            previous = _summarize_runs(
                _fetch_runs(
                    self.database, project_id, None, runs[-1]["created_at"], limit
                )
            )

        schedule_params: list[Any] = [project_id, from_iso or "1970-01-01T00:00:00.000Z"]
        schedule_query = """
            SELECT action, COUNT(*) AS count FROM audit_events
            WHERE project_id = ? AND action IN ('schedule.triggered', 'schedule.skipped')
              AND created_at >= ?
        """
        if to_iso:
            schedule_query += " AND created_at <= ?"
            schedule_params.append(to_iso)
        schedule_query += " GROUP BY action"
        schedule_rows = self.database.execute(
            schedule_query, tuple(schedule_params)
        ).fetchall()
        triggered = 0
        skipped = 0
        for row in schedule_rows:
            if row[0] == "schedule.triggered":
                triggered = int(row[1])
            elif row[0] == "schedule.skipped":
                skipped = int(row[1])
        schedule_total = triggered + skipped

        return {
            "summary": _summarize_runs(runs),
            "previous": previous,
            "trend": sorted(trend.values(), key=lambda item: item["date"])[-30:],
            "failureCategories": sorted(
                (
                    {"category": category, "count": count, "dimension": category_by}
                    for category, count in categories.items()
                ),
                key=lambda item: item["count"],
                reverse=True,
            ),
            "slowSteps": sorted(
                (
                    {
                        "stepId": item["stepId"],
                        "title": item["title"],
                        "count": item["count"],
                        "averageMs": round(item["totalMs"] / item["count"]),
                        "maxMs": item["maxMs"],
                    }
                    for item in slow_steps.values()
                ),
                key=lambda item: item["averageMs"],
                reverse=True,
            )[:20],
            "elementImpact": sorted(
                (
                    {
                        "elementId": item["elementId"],
                        "name": item["name"],
                        "runCount": item["runCount"],
                        "flowCount": len(item["flowCount"]),
                        "failedRuns": item["failedRuns"],
                        "lastUsedAt": item["lastUsedAt"],
                    }
                    for item in elements.values()
                ),
                key=lambda item: item["runCount"],
                reverse=True,
            )[:100],
            "runDurations": sorted(
                (
                    {
                        "date": item["date"],
                        "averageMs": round(item["totalMs"] / item["count"]),
                        "count": item["count"],
                    }
                    for item in durations.values()
                ),
                key=lambda item: item["date"],
            )[-30:],
            "scheduleHealth": {
                "triggered": triggered,
                "skipped": skipped,
                "successRate": round((triggered / schedule_total) * 100)
                if schedule_total
                else 0,
            },
        }

    def process_due_schedules(self) -> None:
        from .http import PlatformError

        rows = self.database.execute(
            """
            SELECT id, project_id, revision_id, environment_id,
                   dataset_version_id, cron_expression, timezone, next_run_at
            FROM schedules
            WHERE enabled = 1 AND archived_at IS NULL AND next_run_at <= ?
              AND project_id NOT IN (
                SELECT id FROM platform_projects WHERE archived_at IS NOT NULL
              )
            ORDER BY next_run_at ASC LIMIT 20
            """,
            (now(),),
        ).fetchall()
        for schedule in rows:
            attempted_at = now()
            try:
                queued = self.queue_published_runs(
                    {
                        "projectId": schedule[1],
                        "revisionId": schedule[2],
                        "environmentId": schedule[3],
                        "datasetVersionId": schedule[4],
                        "createdBy": f"schedule:{schedule[0]}",
                        "source": "schedule",
                        "dispatchKey": f"schedule:{schedule[0]}:{schedule[7]}",
                    }
                )
                self.database.execute(
                    """
                    UPDATE schedules
                    SET last_run_at = ?, next_run_at = ?, updated_at = ?
                    WHERE id = ?
                    """,
                    (attempted_at, next_cron_time(schedule[5], schedule[6]), now(), schedule[0]),
                )
                project = self.project_for(schedule[1])
                self.audit(
                    project["workspace_id"],
                    {"type": "system", "id": f"schedule:{schedule[0]}"},
                    "schedule.triggered",
                    {"type": "schedule", "id": schedule[0]},
                    {"runIds": queued["runIds"]},
                    schedule[1],
                )
            except Exception as exc:
                try:
                    next_run_at = next_cron_time(schedule[5], schedule[6])
                except Exception:
                    next_run_at = _iso_add_seconds(60)
                self.database.execute(
                    """
                    UPDATE schedules SET next_run_at = ?, updated_at = ?
                    WHERE id = ?
                    """,
                    (next_run_at, now(), schedule[0]),
                )
                project = self.project_for(schedule[1])
                self.audit(
                    project["workspace_id"],
                    {"type": "system", "id": f"schedule:{schedule[0]}"},
                    "schedule.skipped",
                    {"type": "schedule", "id": schedule[0]},
                    {
                        "error": (
                            exc.code
                            if isinstance(exc, PlatformError)
                            else "SCHEDULE_TRIGGER_FAILED"
                        )
                    },
                    schedule[1],
                )


def _now_ms() -> float:
    return datetime.now(timezone.utc).timestamp() * 1000


def _read_xlsx(content: bytes) -> list[list[Any]]:
    from datetime import date, datetime

    import openpyxl

    workbook = openpyxl.load_workbook(
        io.BytesIO(content), read_only=True, data_only=True
    )
    sheet = workbook.active
    rows: list[list[Any]] = []
    for values in sheet.iter_rows(values_only=True):
        row: list[Any] = []
        for value in values:
            if isinstance(value, datetime):
                row.append(value.strftime("%Y-%m-%dT%H:%M:%S.%f")[:-3] + "Z")
            elif isinstance(value, date):
                row.append(value.isoformat())
            elif value is None:
                row.append("")
            else:
                row.append(value)
        rows.append(row)
    workbook.close()
    return rows


def _parse_iso(value: str) -> datetime:
    return datetime.fromisoformat(value.replace("Z", "+00:00"))


def _iso_from_ms(value: float) -> str:
    return (
        datetime.fromtimestamp(value / 1000, timezone.utc).strftime(
            "%Y-%m-%dT%H:%M:%S.%f"
        )[:-3]
        + "Z"
    )


def _iso_ms(value: str | None) -> float | None:
    if not value:
        return None
    try:
        parsed = _parse_iso(value)
    except ValueError:
        return None
    if parsed.tzinfo is None:
        parsed = parsed.replace(tzinfo=timezone.utc)
    return parsed.timestamp() * 1000


def _period_key(iso: str, period: str) -> str:
    if period == "day":
        return iso[:10]
    date_value = _parse_iso(iso).date()
    year, week, _ = date_value.isocalendar()
    return f"{year}-W{week:02d}"


def _fetch_runs(
    database: sqlite3.Connection,
    project_id: str,
    from_iso: str | None,
    to_iso: str | None,
    limit: int = 2000,
) -> list[dict[str, Any]]:
    query = """
        SELECT id, revision_id, status, snapshot, created_at
        FROM platform_runs WHERE project_id = ?
    """
    params: list[Any] = [project_id]
    if from_iso:
        query += " AND created_at >= ?"
        params.append(from_iso)
    if to_iso:
        query += " AND created_at <= ?"
        params.append(to_iso)
    query += " ORDER BY created_at DESC LIMIT ?"
    params.append(limit)
    return [
        {
            "id": row[0],
            "revision_id": row[1],
            "status": row[2],
            "snapshot": row[3],
            "created_at": row[4],
        }
        for row in database.execute(query, tuple(params)).fetchall()
    ]


def _summarize_runs(run_rows: list[dict[str, Any]]) -> dict[str, int]:
    terminal = [
        run for run in run_rows if run["status"] in ("success", "failed", "canceled")
    ]
    failed_runs = sum(1 for run in run_rows if run["status"] == "failed")
    canceled_runs = sum(1 for run in run_rows if run["status"] == "canceled")
    success_runs = sum(1 for run in terminal if run["status"] == "success")
    return {
        "totalRuns": len(run_rows),
        "successRate": round((success_runs / len(terminal)) * 100) if terminal else 0,
        "failedRuns": failed_runs,
        "canceledRuns": canceled_runs,
        "failedRate": round((failed_runs / len(run_rows)) * 100) if run_rows else 0,
        "canceledRate": round((canceled_runs / len(run_rows)) * 100) if run_rows else 0,
    }
